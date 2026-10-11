"""
doc_render.py
=============
번역 결과 DOCX 렌더러 + DOCX->PDF 변환 (v4).

서식 (사용자 형식 샘플 Biddle1998_..._번역.docx 의 스타일 값을 그대로 따름, 2026-10-11)
  - A4, 여백 2.5cm, 본문 맑은 고딕 11pt, 줄간격 1.5, 문단 뒤 7pt, 양쪽 정렬
  - 제목 18pt 굵게 가운데 / 저자 줄 10pt 굵게 가운데 / 서지 머리줄 회색 8.5pt 우측
  - ## 절 14pt 굵게(앞 16pt) / ### 소절 12pt 굵게(앞 12pt) — 구분선 없음
  - 인용 블록 10.5pt 좌우 1cm 들여쓰기 / 그림 캡션 10pt 굵게 가운데
  - [원서 p. N] 쪽 마커: 회색 8.5pt (독립 줄이면 PageMarker 단락 스타일, 우측 정렬, 다음 단락과 함께;
    문단 중간이면 PageMarkerChar 문자 스타일)
  - 각주: 실제 Word 각주 (footnotes.xml), 9pt
  - 그림: 본문 폭의 92% 이내, 가운데

PDF 는 이 DOCX 를 Word(COM) -> LibreOffice 순으로 변환한다. 둘 다 실패하면 호출 측이 ReportLab 으로 폴백.
"""

import os
import re
import copy
import shutil
import subprocess
import tempfile
from pathlib import Path

FONT = "맑은 고딕"
BODY_PT = 11
NOTE_PT = 9
MARKER_PT = 8.5
MARKER_RGB = (0x80, 0x80, 0x80)
LINE_SPACING = 1.5
SPACE_AFTER_PT = 7
MARGIN_CM = 2.5
IMG_MAX_RATIO = 0.92

PAGE_TOKEN_RE = re.compile(r"\[\[\s*PAGE\s*:\s*([0-9]{1,4}|[ivxlcdmIVXLCDM]{1,8})\s*\]\]")
FNREF_RE = re.compile(r"\[\^([A-Za-z0-9_\-]+)\](?!:)")

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
FOOTNOTES_CT = "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"
FOOTNOTES_RT = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes"


def marker_label(n: str) -> str:
    return f"[원서 p. {n}]"


# ──────────────────────────────────────────────
# 스타일
# ──────────────────────────────────────────────
def _set_font(style_or_run_font_owner, size_pt=None, bold=None, italic=None, color=None):
    from docx.shared import Pt, RGBColor
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    f = style_or_run_font_owner.font
    f.name = FONT
    if size_pt is not None:
        f.size = Pt(size_pt)
    if bold is not None:
        f.bold = bold
    if italic is not None:
        f.italic = italic
    if color is not None:
        f.color.rgb = RGBColor(*color)
    el = style_or_run_font_owner.element if hasattr(style_or_run_font_owner, "element") else style_or_run_font_owner._r
    rPr = el.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    for k in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rFonts.set(qn(k), FONT)


def _para_border_top(paragraph_format_owner_el, sz=12):
    """단락(또는 단락 스타일)의 pPr 에 상단 굵은 선."""
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    pPr = paragraph_format_owner_el.get_or_add_pPr()
    pBdr = pPr.find(qn("w:pBdr"))
    if pBdr is None:
        pBdr = OxmlElement("w:pBdr")
        pPr.append(pBdr)
    top = OxmlElement("w:top")
    top.set(qn("w:val"), "single")
    top.set(qn("w:sz"), str(sz))
    top.set(qn("w:space"), "6")
    top.set(qn("w:color"), "000000")
    pBdr.append(top)


def _set_doc_defaults(doc):
    """테마 글꼴(Calibri / 中文 대체 글꼴) 대신 문서 기본 글꼴을 맑은 고딕으로, 언어는 ko-KR."""
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    styles_el = doc.styles.element
    dd = styles_el.find(qn("w:docDefaults"))
    if dd is None:
        return
    rpd = dd.find(qn("w:rPrDefault"))
    if rpd is None:
        rpd = OxmlElement("w:rPrDefault"); dd.insert(0, rpd)
    rpr = rpd.find(qn("w:rPr"))
    if rpr is None:
        rpr = OxmlElement("w:rPr"); rpd.append(rpr)
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.insert(0, rf)
    for k in list(rf.attrib):
        del rf.attrib[k]
    for k in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(k), FONT)
    lang = rpr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang"); rpr.append(lang)
    lang.set(qn("w:val"), "en-US")
    lang.set(qn("w:eastAsia"), "ko-KR")


def _setup_styles(doc):
    from docx.shared import Pt
    from docx.enum.style import WD_STYLE_TYPE
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    _set_doc_defaults(doc)

    st = doc.styles
    normal = st["Normal"]
    _set_font(normal, BODY_PT)
    pf = normal.paragraph_format
    pf.line_spacing = LINE_SPACING
    pf.space_after = Pt(SPACE_AFTER_PT)
    pf.space_before = Pt(0)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    def heading(name, size, before, after, border=False, center=False):
        s = st[name]
        _set_font(s, size, bold=True, italic=False, color=(0, 0, 0))
        f = s.paragraph_format
        f.space_before = Pt(before)
        f.space_after = Pt(after)
        f.keep_with_next = True
        f.line_spacing = 1.2
        f.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
        if border:
            _para_border_top(s.element)
        return s

    t = heading("Title", 18, 0, 8, center=True)
    from docx.oxml.ns import qn as _qn
    _tp = t.element.find(_qn("w:pPr"))
    if _tp is not None:                       # 기본 템플릿 Title 의 파란 하단선 제거
        for bd in _tp.findall(_qn("w:pBdr")):
            _tp.remove(bd)
    heading("Heading 1", 14, 16, 8)                   # ## 절
    heading("Heading 2", 12, 12, 8)                   # ### 소절
    heading("Heading 3", 11, 10, 6)                   # #### 이하

    for nm in ("Subtitle",):
        try:
            s = st[nm]
            _set_font(s, 12, bold=True, italic=False, color=(0x33, 0x33, 0x33))
            s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            s.paragraph_format.space_after = Pt(6)
        except KeyError:
            pass

    # 쪽 마커 (단락/문자)
    if "PageMarker" not in [s.name for s in st]:
        pm = st.add_style("PageMarker", WD_STYLE_TYPE.PARAGRAPH)
        pm.base_style = normal
        _set_font(pm, MARKER_PT, color=MARKER_RGB)
        pm.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        pm.paragraph_format.space_before = Pt(2)
        pm.paragraph_format.space_after = Pt(2)
        pm.paragraph_format.line_spacing = 1.0
        pm.paragraph_format.keep_with_next = True
    if "PageMarkerChar" not in [s.name for s in st]:
        pc = st.add_style("PageMarkerChar", WD_STYLE_TYPE.CHARACTER)
        _set_font(pc, MARKER_PT, color=MARKER_RGB)

    # 저자 줄 / 그림 캡션 / 서지 머리줄 / 인용 블록
    def _pstyle(name, size, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER, before=0, after=7,
                color=None, line=None, left=None, right=None, keep=False):
        from docx.shared import Cm
        if name not in [x.name for x in st]:
            ps = st.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            ps.base_style = normal
        ps = st[name]
        _set_font(ps, size, bold=bold, color=color)
        f = ps.paragraph_format
        f.alignment = align
        f.space_before = Pt(before)
        f.space_after = Pt(after)
        if line: f.line_spacing = line
        if left is not None: f.left_indent = Cm(left)
        if right is not None: f.right_indent = Cm(right)
        if keep: f.keep_with_next = True
        return ps
    _pstyle("AuthorLine", 10, bold=True, after=10, keep=True)
    _pstyle("FigCaption", 10, bold=True, after=14)
    _pstyle("SourceLine", MARKER_PT, align=WD_ALIGN_PARAGRAPH.RIGHT, after=2, color=MARKER_RGB, line=1.2)
    _pstyle("BlockQuote", 10.5, align=WD_ALIGN_PARAGRAPH.JUSTIFY, before=5, after=5, line=1.4,
            left=1.0, right=1.0)

    # 각주 스타일
    names = [s.name for s in st]
    if "Footnote Text" not in names:
        ft = st.add_style("Footnote Text", WD_STYLE_TYPE.PARAGRAPH)
        ft.base_style = normal
    ft = st["Footnote Text"]
    _set_font(ft, NOTE_PT)
    ft.paragraph_format.line_spacing = 1.15
    ft.paragraph_format.space_after = Pt(3)
    ft.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if "Footnote Reference" not in names:
        fr = st.add_style("Footnote Reference", WD_STYLE_TYPE.CHARACTER)
        fr.font.superscript = True

    for nm in ("Quote", "List Paragraph"):
        try:
            _set_font(st[nm], BODY_PT, color=(0, 0, 0))
        except KeyError:
            pass


# ──────────────────────────────────────────────
# 각주 파트
# ──────────────────────────────────────────────
class _Footnotes:
    def __init__(self, doc):
        from docx.opc.part import Part
        from docx.opc.packuri import PackURI
        self.doc = doc
        self.next_id = 1
        self.items = []          # (id, text)
        self._Part, self._PackURI = Part, PackURI

    def add(self, text: str) -> int:
        fid = self.next_id
        self.next_id += 1
        self.items.append((fid, text))
        return fid

    def finalize(self, add_runs_fn):
        """footnotes.xml 생성 후 패키지에 연결."""
        if not self.items:
            return
        from docx.oxml import parse_xml
        from docx.oxml.ns import nsdecls
        xml = (f'<w:footnotes {nsdecls("w", "r")}>'
               '<w:footnote w:type="separator" w:id="-1"><w:p><w:pPr><w:spacing w:after="0" w:line="240" w:lineRule="auto"/></w:pPr>'
               '<w:r><w:separator/></w:r></w:p></w:footnote>'
               '<w:footnote w:type="continuationSeparator" w:id="0"><w:p><w:pPr><w:spacing w:after="0" w:line="240" w:lineRule="auto"/></w:pPr>'
               '<w:r><w:continuationSeparator/></w:r></w:p></w:footnote>'
               '</w:footnotes>')
        root = parse_xml(xml)
        # 각주 단락은 임시 단락을 만들어 run 구성을 재사용한 뒤 옮긴다
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        tmp_doc_body = self.doc.add_paragraph()
        ft_style_id = self.doc.styles["Footnote Text"].style_id
        fr_style_id = self.doc.styles["Footnote Reference"].style_id
        for fid, text in self.items:
            fn = OxmlElement("w:footnote")
            fn.set(qn("w:id"), str(fid))
            p = copy.deepcopy(tmp_doc_body._p)
            for child in list(p):
                p.remove(child)
            from docx.text.paragraph import Paragraph
            para = Paragraph(p, tmp_doc_body._parent)
            para.style = self.doc.styles["Footnote Text"]
            r = OxmlElement("w:r")
            rPr = OxmlElement("w:rPr")
            rs = OxmlElement("w:rStyle"); rs.set(qn("w:val"), fr_style_id)
            va = OxmlElement("w:vertAlign"); va.set(qn("w:val"), "superscript")
            rPr.append(rs); rPr.append(va); r.append(rPr)
            r.append(OxmlElement("w:footnoteRef"))
            p.append(r)
            sp = OxmlElement("w:r"); t = OxmlElement("w:t"); t.text = " "
            t.set(qn("xml:space"), "preserve"); sp.append(t); p.append(sp)
            add_runs_fn(para, text, NOTE_PT)
            fn.append(p)
            root.append(fn)
        tmp_doc_body._p.getparent().remove(tmp_doc_body._p)

        from lxml import etree
        blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
        part = self._Part(self._PackURI("/word/footnotes.xml"), FOOTNOTES_CT, blob, self.doc.part.package)
        self.doc.part.relate_to(part, FOOTNOTES_RT)


def _add_footnote_ref(paragraph, fid: int, style_id: str):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    r = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    rs = OxmlElement("w:rStyle"); rs.set(qn("w:val"), style_id)
    va = OxmlElement("w:vertAlign"); va.set(qn("w:val"), "superscript")
    rPr.append(rs); rPr.append(va); r.append(rPr)
    ref = OxmlElement("w:footnoteReference"); ref.set(qn("w:id"), str(fid))
    r.append(ref)
    paragraph._p.append(r)


# ──────────────────────────────────────────────
# 런 구성 (굵게/기울임 + 인라인 쪽 마커 + 각주 참조)
# ──────────────────────────────────────────────
def _add_md_runs(paragraph, text: str, size_pt=None, italic=False):
    pat = r"(\*\*\*.+?\*\*\*|\*\*.+?\*\*|\*.+?\*)"
    for part in re.split(pat, text):
        if not part:
            continue
        b = i = False
        if part.startswith("***") and part.endswith("***") and len(part) > 6:
            part, b, i = part[3:-3], True, True
        elif part.startswith("**") and part.endswith("**") and len(part) > 4:
            part, b = part[2:-2], True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            part, i = part[1:-1], True
        run = paragraph.add_run(part)
        _set_font(run, size_pt, bold=b or None, italic=(i or italic) or None)


def _add_rich(paragraph, text: str, notes: dict, fn: "_Footnotes", size_pt=None, italic=False):
    """text 안의 [[PAGE:N]] / [^id] 를 처리하며 run 추가."""
    doc_styles = paragraph.part.document.styles if hasattr(paragraph.part, "document") else None
    pos = 0
    tok_re = re.compile(PAGE_TOKEN_RE.pattern + r"|" + FNREF_RE.pattern)
    for m in tok_re.finditer(text):
        pre = text[pos:m.start()]
        if pre:
            _add_md_runs(paragraph, pre, size_pt, italic)
        if m.group(1) is not None:          # 인라인 쪽 마커
            lead = "" if (pre.endswith(" ") or not pre) else " "
            run = paragraph.add_run(lead + marker_label(m.group(1)) + " ")
            run.style = doc_styles["PageMarkerChar"] if doc_styles else None
        else:                               # 각주 참조
            nid = m.group(2)
            ntext = notes.get(nid)
            if ntext and fn is not None:
                fid = fn.add(ntext)
                _add_footnote_ref(paragraph, fid, doc_styles["Footnote Reference"].style_id)
        pos = m.end()
    rest = text[pos:]
    if rest:
        _add_md_runs(paragraph, rest, size_pt, italic)
    # 런 사이 이중 공백 정리
    for r in paragraph.runs:
        if r.text and "  " in r.text:
            r.text = re.sub(r" {2,}", " ", r.text)


def build_docx(blocks, notes: dict, output_path: str, assets: dict = None, image_size_fn=None,
               source_line: str = ""):
    """blocks: (kind, payload) 시퀀스 (paper_utils._iter_blocks). notes: {id: 번역문}"""
    from docx import Document
    from docx.shared import Pt, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    assets = assets or {}
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    for side in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(sec, side, Cm(MARGIN_CM))
    _setup_styles(doc)
    text_w_cm = 21.0 - 2 * MARGIN_CM
    fn = _Footnotes(doc)

    # 쪽번호 바닥글 (- N -)
    try:
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        fp = sec.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        def _fld(code):
            r = fp.add_run()
            b = OxmlElement("w:fldChar"); b.set(qn("w:fldCharType"), "begin"); r._r.append(b)
            r2 = fp.add_run(); it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve")
            it.text = code; r2._r.append(it)
            r3 = fp.add_run(); e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end"); r3._r.append(e)
        r0 = fp.add_run("- "); _set_font(r0, 9)
        _fld(" PAGE ")
        r1 = fp.add_run(" -"); _set_font(r1, 9)
    except Exception:
        pass

    source_done = False

    def _source_line():
        p = doc.add_paragraph(style="SourceLine")
        _add_md_runs(p, source_line)

    for kind, payload in blocks:
        if source_line and not source_done and kind not in ("title", "author"):
            _source_line(); source_done = True
        if kind == "page":
            p = doc.add_paragraph(marker_label(payload), style="PageMarker")
            continue
        if kind == "image":
            a = assets.get(payload)
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            if a and Path(a["path"]).exists():
                w_cm = image_size_fn(a, max_w_cm=text_w_cm * IMG_MAX_RATIO)[0] if image_size_fn else text_w_cm * 0.8
                try:
                    p.add_run().add_picture(a["path"], width=Cm(w_cm))
                except Exception as e:
                    _add_md_runs(p, f"[이미지 삽입 실패: {payload}: {e}]")
            else:
                _add_md_runs(p, f"[이미지 누락: {payload}]")
            continue
        if kind == "table":
            rows = [r.strip().strip("|").split("|") for k, r in enumerate(payload) if k != 1]
            ncol = max(len(r) for r in rows)
            tbl = doc.add_table(rows=len(rows), cols=ncol)
            try:
                tbl.style = doc.styles["Table Grid"]
            except KeyError:
                pass
            for ri, r in enumerate(rows):
                for ci in range(ncol):
                    cell = tbl.cell(ri, ci)
                    cp = cell.paragraphs[0]
                    cp.paragraph_format.space_after = Pt(0)
                    cp.paragraph_format.line_spacing = 1.15
                    txt = r[ci].strip() if ci < len(r) else ""
                    _add_rich(cp, f"**{txt}**" if (ri == 0 and txt) else txt, notes, fn, size_pt=9.5)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
            continue
        if kind == "title":
            p = doc.add_paragraph(style="Title")
            _add_rich(p, payload, notes, fn)
            continue
        if kind == "subtitle":
            p = doc.add_paragraph(style="Subtitle")
            _add_rich(p, payload, notes, fn)
            continue
        if kind == "caption":
            p = doc.add_paragraph(style="FigCaption")
            _add_rich(p, payload, notes, fn)
            continue
        if kind == "heading":
            level, txt = payload
            style = "Heading 1" if level <= 2 else ("Heading 2" if level == 3 else "Heading 3")
            p = doc.add_paragraph(style=style)
            _add_rich(p, txt, notes, fn)
            continue
        if kind == "list":
            p = doc.add_paragraph(style="List Paragraph")
            p.paragraph_format.space_after = Pt(3)
            _add_rich(p, "• " + payload, notes, fn)
            continue
        if kind == "quote":
            p = doc.add_paragraph(style="BlockQuote")
            _add_rich(p, payload, notes, fn)
            continue
        if kind == "author":
            p = doc.add_paragraph(style="AuthorLine")
            _add_rich(p, payload, notes, fn)
            continue
        p = doc.add_paragraph()
        _add_rich(p, payload, notes, fn)

    if source_line and not source_done:
        _source_line()
    fn.finalize(lambda para, text, size: _add_md_runs(para, text, size))
    doc.save(output_path)
    return len(fn.items)


# ──────────────────────────────────────────────
# DOCX -> PDF
# ──────────────────────────────────────────────
def _word_to_pdf(docx_path: str, pdf_path: str) -> bool:
    if os.name != "nt":
        return False
    try:
        import pythoncom
        import win32com.client
    except ImportError:
        return False
    word = None
    try:
        pythoncom.CoInitialize()
        word = win32com.client.DispatchEx("Word.Application")
        word.Visible = False
        word.DisplayAlerts = 0
        d = word.Documents.Open(str(Path(docx_path).resolve()), ReadOnly=True,
                                AddToRecentFiles=False, ConfirmConversions=False)
        d.ExportAsFixedFormat(str(Path(pdf_path).resolve()), 17)   # wdExportFormatPDF
        d.Close(False)
        return Path(pdf_path).exists()
    except Exception as e:
        print(f"    (Word 변환 실패: {e})", flush=True)
        return False
    finally:
        try:
            if word is not None:
                word.Quit()
        except Exception:
            pass
        try:
            pythoncom.CoUninitialize()
        except Exception:
            pass


def _soffice_path():
    for c in ("soffice", "soffice.exe"):
        p = shutil.which(c)
        if p:
            return p
    for c in (r"C:\Program Files\LibreOffice\program\soffice.exe",
              r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"):
        if os.path.exists(c):
            return c
    return None


def _soffice_to_pdf(docx_path: str, pdf_path: str) -> bool:
    exe = _soffice_path()
    if not exe:
        return False
    out_dir = tempfile.mkdtemp(prefix="tr_pdf_")
    try:
        subprocess.run([exe, "--headless", "--convert-to", "pdf", "--outdir", out_dir, str(docx_path)],
                       check=True, timeout=300, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        produced = Path(out_dir) / (Path(docx_path).stem + ".pdf")
        if produced.exists():
            shutil.move(str(produced), pdf_path)
            return True
    except Exception as e:
        print(f"    (LibreOffice 변환 실패: {e})", flush=True)
    finally:
        shutil.rmtree(out_dir, ignore_errors=True)
    return False


def docx_to_pdf(docx_path: str, pdf_path: str) -> str:
    """성공 시 사용한 엔진 이름('Word'|'LibreOffice'), 실패 시 ''."""
    if _word_to_pdf(docx_path, pdf_path):
        return "Word"
    if _soffice_to_pdf(docx_path, pdf_path):
        return "LibreOffice"
    return ""
