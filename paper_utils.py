"""
paper_utils.py
==============
논문 번역 파이프라인 공통 유틸리티 (v4).
translate_paper.py / batch_submit.py / batch_download.py 에서 공유.

v4 변경점 (2026-10, '리딩 번역 작업 지침' 형식 반영)
  - 모델: claude-fable-5-1 / claude-opus-5-5 / claude-sonnet-5-5(기본) / claude-haiku-4-5, 가격표 갱신
  - 원서 쪽 마커 [원서 p. N] (인쇄 쪽번호 자동 감지, 비전 모드는 모델이 쪽번호를 읽음)
  - 각주 3모드: strip(삭제) / commentary(논평성 각주만 Word 각주로 번역, 서지는 영문 유지) / all
  - 사용자 용어집(user_glossary.txt) 우선 적용, 번역투 지양 스타일
  - DOCX 서식: A4·여백 2.5cm·맑은 고딕 11pt·줄간격 1.5·문단 뒤 7pt·양쪽정렬, 실제 각주 9pt
  - PDF: DOCX 를 Word(COM)로 변환 -> LibreOffice -> ReportLab 순 폴백, 'docx+pdf' 동시 출력
  - 출력 파일명: <원본>_pp<시작>-<끝>_한국어번역.<ext>
  - QA: 쪽 마커 순차, 각주 짝, 쪽별 한글 글자수/원문 단어수 비율 (요약본 의심 청크 자동 재번역)

v3 변경점
  - PDF 추출을 pdf_extract.py 로 분리: 단락 보존 + 페이지 경계 병합 + 도식/표/수식 이미지화
  - 청크 확대(6000자) + 문장 단위 안전 분할 -> 문장 중간 절단 없음
  - 용어집(glossary) 사전 추출 패스 + 직전 청크 문맥 제공 -> 용어 일관성/문맥 연속성
  - 최신 모델(claude-fable-5 / claude-opus-5 / claude-sonnet-5) 가격표, effort 옵션
  - 프롬프트 캐싱(system 블록 cache_control)
  - DOCX/PDF/MD 저장 시 [[IMG:id]] 마커를 원본 이미지로 치환
"""

import os
import sys
import re
import json
import time
from collections import Counter
from pathlib import Path

BASE_DIR = Path(__file__).parent

# ──────────────────────────────────────────────
# 설정
# ──────────────────────────────────────────────
def _load_api_key():
    env = os.environ.get("ANTHROPIC_API_KEY")
    if env:
        return env
    cfg = BASE_DIR / "api_key.cfg"
    if cfg.exists():
        try:
            k = cfg.read_text(encoding="utf-8").strip()
            if k:
                return k
        except Exception:
            pass
    return None


API_KEY = _load_api_key()


def _load_model() -> str:
    cfg = BASE_DIR / "model.cfg"
    if cfg.exists():
        try:
            m = cfg.read_text(encoding="utf-8").strip()
            if m:
                return m
        except Exception:
            pass
    return DEFAULT_MODEL


DEFAULT_MODEL = "claude-sonnet-5-5"
MODEL = _load_model()


def _load_extract_cfg() -> dict:
    """extract.cfg (key=value 한 줄씩)

        footnotes=commentary      strip(삭제) | commentary(논평성만 번역) | all(전체 번역)
        strip_references=true     References 이후 절단
        output_format=docx        docx | pdf | docx+pdf | md
        effort=low                off | low | medium | high | xhigh  (thinking 강도)
        figures=true              도식/표/수식 이미지 추출
        glossary=true             번역 전 용어집 추출 패스 (API 1회 추가)
        page_markers=true         [원서 p. N] 쪽 마커
        first_page=auto           원서 첫 쪽 번호 (auto = 인쇄 쪽번호 자동 감지)
        glossary_file=none        사용자 용어집: glossaries/ 안의 파일 이름 (none = 사용 안 함)
        (하위 호환) strip_footnotes=true/false -> footnotes=strip/all
    """
    bool_defaults = {"strip_references": True, "figures": True, "glossary": True,
                     "page_markers": True}
    str_defaults  = {"output_format": "docx", "effort": "low", "footnotes": "",
                     "first_page": "auto", "glossary_file": ""}
    out = {**bool_defaults, **str_defaults}
    legacy_strip = None
    cfg = BASE_DIR / "extract.cfg"
    if not cfg.exists():
        return out
    try:
        for line in cfg.read_text(encoding="utf-8").splitlines():
            line = line.strip()
            if not line or line.startswith("#") or "=" not in line:
                continue
            k, v = line.split("=", 1)
            k = k.strip().lower(); v = v.strip().lower()
            if k in bool_defaults:
                out[k] = (v == "true")
            elif k == "strip_footnotes":
                legacy_strip = (v == "true")
            elif k == "footnotes" and v in ("strip", "commentary", "all"):
                out[k] = v
            elif k == "output_format" and v in ("docx", "pdf", "docx+pdf", "md"):
                out[k] = v
            elif k == "effort" and v in ("off", "low", "medium", "high", "xhigh"):
                out[k] = v
            elif k == "glossary_file":
                out[k] = line.split("=", 1)[1].strip()      # 대소문자 보존
            elif k == "first_page":
                out[k] = v if (v == "auto" or v.isdigit()) else "auto"
    except Exception:
        pass
    if not out["footnotes"]:
        out["footnotes"] = "commentary" if legacy_strip is None else ("strip" if legacy_strip else "all")
    return out


_EXTRACT_CFG     = _load_extract_cfg()
FOOTNOTE_MODE    = _EXTRACT_CFG["footnotes"]
STRIP_FOOTNOTES  = FOOTNOTE_MODE == "strip"
PAGE_MARKERS     = _EXTRACT_CFG["page_markers"]
FIRST_PAGE       = int(_EXTRACT_CFG["first_page"]) if _EXTRACT_CFG["first_page"].isdigit() else None
STRIP_REFERENCES = _EXTRACT_CFG["strip_references"]
OUTPUT_FORMAT    = _EXTRACT_CFG["output_format"]
EFFORT           = _EXTRACT_CFG["effort"]
EXTRACT_FIGURES  = _EXTRACT_CFG["figures"]
USE_GLOSSARY     = _EXTRACT_CFG["glossary"]

MAX_CHUNK = 6000   # 청크당 최대 문자 수 (문장 경계에서만 분할)
MIN_CHUNK = 20

# 요청 1건당 출력 토큰 상한. 6000자 영문 -> 한국어 약 3500~4500 토큰 + thinking 여유.
MAX_OUTPUT_TOKENS = 16000

# 직전 청크 문맥 길이 (원문 / 번역문 꼬리)
CONTEXT_SRC_CHARS = 700
CONTEXT_TR_CHARS  = 500

ASSETS_ROOT = BASE_DIR / "assets"

MARKER_RE = re.compile(r"\[\[\s*IMG\s*:\s*([A-Za-z0-9_\-]+)\s*\]\]")
PAGE_RE   = re.compile(r"\[\[\s*PAGE\s*:\s*([0-9]{1,4}|[ivxlcdmIVXLCDM]{1,8})\s*\]\]")
FNREF_RE  = re.compile(r"\[\^([A-Za-z0-9_\-]+)\](?!:)")
FNDEF_RE  = re.compile(r"^[ \t]*\[\^([A-Za-z0-9_\-]+)\]:[ \t]*(.*)$", re.MULTILINE)
GLOSSARY_DIR = BASE_DIR / "glossaries"


def _resolve_user_glossary(name: str):
    """glossary_file 설정 -> 경로. 비어 있으면 예전 위치(user_glossary.txt)를 하위 호환으로 사용."""
    if not name:
        legacy = BASE_DIR / "user_glossary.txt"
        return legacy if legacy.exists() else None
    if name.lower() in ("none", "off", "false"):
        return None
    for cand in (GLOSSARY_DIR / name, BASE_DIR / name, Path(name)):
        if cand.exists():
            return cand
    return None


USER_GLOSSARY_PATH = _resolve_user_glossary(_EXTRACT_CFG["glossary_file"])


# ──────────────────────────────────────────────
# thinking / effort 설정 (모델별)
# ──────────────────────────────────────────────
def get_thinking_kwargs(model: str, effort: str = None) -> dict:
    """모델 + effort 옵션에 따른 messages.create()/batch 추가 kwargs.

    (platform.claude.com/docs 'Thinking' 표 기준, 2026-10 확인)
    - fable-5-1 / fable-5 / mythos / opus-5-5 : adaptive 고정. disabled 는 400 -> effort 로만 조절
    - opus-5      : adaptive 기본. disabled 는 effort high 이하에서만 허용
    - sonnet-5-5  : adaptive 기본. disabled 는 effort high 이하에서 허용
    - sonnet-5    : adaptive 기본. disabled 허용
    - opus-4-8 / 4-7 : adaptive, 기본 Off. disabled 허용
    - haiku-4-5 등 : extended thinking(budget) 방식, 기본 Off -> 손대지 않음
    effort 값(off/low/medium/high/xhigh)은 output_config.effort 로 전달.
    effort = off : thinking 끔 (끌 수 없는 모델은 low 로 대체)
    """
    m = (model or "").lower()
    e = (effort or EFFORT or "low").lower()
    always_on = any(k in m for k in ("fable", "mythos", "opus-5-5"))
    supports_effort = always_on or any(k in m for k in ("opus-5", "sonnet-5", "opus-4-8", "opus-4-7"))
    if not supports_effort:
        return {}
    if e == "off":
        if always_on:
            return {"thinking": {"type": "adaptive"}, "output_config": {"effort": "low"}}
        return {"thinking": {"type": "disabled"}}
    return {"thinking": {"type": "adaptive"}, "output_config": {"effort": e}}


# ──────────────────────────────────────────────
# API 응답 파싱 / 재시도
# ──────────────────────────────────────────────
def extract_response_text(content) -> str:
    if not content:
        return ""
    return "".join(
        block.text for block in content
        if getattr(block, "type", None) == "text" and hasattr(block, "text")
    )


def call_with_retry(fn, tries: int = 4, base_wait: float = 3.0):
    """rate limit / overloaded / 네트워크 오류 시 지수 백오프 재시도."""
    last = None
    for attempt in range(tries):
        try:
            return fn()
        except Exception as e:  # noqa
            last = e
            name = type(e).__name__
            msg = str(e)
            retriable = any(k in name for k in ("RateLimit", "Overloaded", "APIConnection",
                                                 "InternalServer", "APITimeout", "APIStatus"))
            retriable = retriable or any(k in msg for k in ("529", "overloaded", "rate_limit", "429", "500", "503"))
            if not retriable or attempt == tries - 1:
                raise
            wait = base_wait * (2 ** attempt)
            print(f"    [재시도] {name}: {wait:.0f}s 대기 후 재시도 ({attempt+1}/{tries-1})", flush=True)
            time.sleep(wait)
    raise last


# ──────────────────────────────────────────────
# 시스템 프롬프트
# ──────────────────────────────────────────────
SYSTEM_PROMPT = """You are an expert academic translator (English -> Korean) specializing in international relations, security studies, military strategy, political science, and quantitative social science. Your reader is a Korean graduate student who needs a complete, faithful Korean rendering of the source text for private study.

## Output contract
- Translate EVERYTHING in <text_to_translate>. No omissions, no summaries, no paraphrase that drops information. Every sentence, including short transitional sentences, single-line paragraphs, headings, list items, and figure/table captions.
- Output ONLY the Korean translation. No preamble, no commentary, no notes about your choices, no "Human:"/"Assistant:" labels.
- Preserve paragraph structure exactly: one source paragraph -> one target paragraph, separated by a blank line. Keep markdown headings (#, ##, ###) and **bold**/*italic* as in the source, with heading text translated into Korean.
- Lines of the form [[IMG:...]] are image placeholders for figures, tables, or display equations that will be re-inserted later. Copy each placeholder EXACTLY as-is on its own line, at the same position relative to the surrounding text. Never translate, remove, reorder, or modify them.
- Tokens of the form [[PAGE:N]] mark where page N of the original book/journal begins. Copy every one EXACTLY, in the same order. If it stands on its own line, keep it on its own line. If it appears inside a paragraph, place it inside the Korean paragraph at the point that corresponds to the same position in the source (between the corresponding words or clauses, never inside a word).
- Tokens of the form [^n12] are footnote references. Copy every one EXACTLY, attached directly (no space) to the end of the Korean word, clause, or sentence that corresponds to where it is attached in the source. Never drop, renumber, or invent them.
- If <text_to_translate> is empty or contains nothing translatable, output a single blank line.

## Style
- Formal written academic Korean (문어체, '~이다/~하다' 체), but natural and easy to read — avoid translationese (번역투). Do not mirror English syntax: restructure long sentences, split them when needed, prefer active Korean constructions, avoid overusing '~에 의해', '~것이다', '그것', '~를 가지다', '~에 있어서'. Never drop content.
- Prefer clear, widely understood Korean terms over obscure Sino-Korean jargon, so that a general educated reader can follow; use the established Korean academic term when one exists.
- Terminology: use the glossary (if given) consistently; user-specified entries take priority. For key technical terms, theories, and proper nouns, add the English in parentheses at the FIRST occurrence within this section only, e.g. 억지(deterrence), 안보 딜레마(security dilemma). Do not repeat the parenthetical for the same term later in the section. Do not add English for ordinary words.
- Person names: transliterate into Korean and add the original in parentheses at first occurrence, e.g. 케네스 월츠(Kenneth Waltz). Institutions, military units, and treaties: standard Korean name + English in parentheses at first occurrence.
- Keep author-year in-text citations, years, statistics, variable names, and inline math exactly as in the source (e.g. (Fearon 1995), p < 0.05, x ∈ X). Keep italicized variable letters as plain letters.
- Quotations: translate the quoted text into Korean, keeping quotation marks.
- Translate figure/table captions fully (e.g. 'Figure 1. The bargaining range' -> '그림 1. 협상 범위'). Translate GFM/markdown tables cell by cell, keeping the table structure.

## Cleanup rules (source is pre-processed PDF/OCR text)
- Remove plain footnote/endnote callout numbers attached to words or punctuation (e.g. 'explanations."2', 'war.15', '<sup>7</sup>', superscript digits) — but NEVER remove [^n..] tokens. Do NOT remove substantive numbers (years, section numbers, statistics, equation numbers like (3)).
- Fix obvious OCR artifacts (broken hyphenation, stray characters) silently when the intended text is unambiguous.
- Skip ONLY clearly non-content lines: standalone page numbers, running headers, database boilerplate (JSTOR, HeinOnline...), download timestamps, citation-format labels (APA/MLA...), book advertisements. When in doubt, translate.
- If the input is already in Korean, output it unchanged.

## Context
<previous_source> and <previous_translation> (if given) show the END of the preceding section for continuity of terminology, pronouns, and argument. They are context only — do NOT translate or repeat them in your output.
"""

GLOSSARY_SYSTEM = """You are preparing a bilingual terminology glossary to support a consistent English->Korean translation of an academic paper in political science / international relations.

Read the paper text and output a glossary of the terms whose Korean rendering must stay consistent across the whole paper:
- key theoretical concepts and technical terms (e.g. commitment problem, bargaining range, private information, audience costs)
- recurring proper nouns: people, institutions, treaties, places, events, named theories/models
- variables or notation with a conventional Korean reading
- the paper title and the author name(s)

Rules:
- 30 to 90 entries, most important first. Skip ordinary vocabulary.
- Use standard Korean academic renderings used in Korean IR/political science scholarship. For person names give the Korean transliteration.
- Format: one entry per line, exactly `English term = 한국어 번역`. No numbering, no explanations, no headings, nothing else.
- EXCEPTION: the very first line must be `CITATION: <bibliographic citation of this text in Chicago style, as complete as the text itself shows: author(s), "title," journal or book (with editors), volume/issue, publisher, year, pages>`. Write *italics* for journal/book titles. If something is not shown, leave it out.
"""


NOTES_SYSTEM = """You are an expert academic translator (English -> Korean) for political science / security studies. You will receive footnotes from a scholarly text, each as <note id="..." n="...">text</note>.

{mode_rule}

Output format — one line per note, in the same order, nothing else:
[^ID]: (원주 N) 한국어 번역
where ID is the note's id attribute and N its n attribute.

Translation rules:
- Formal written academic Korean, natural and free of translationese; translate every sentence of commentary completely, no summarizing.
- Keep bibliographic citations (author names, titles, journals, publishers, years, page numbers, URLs) in their ORIGINAL English form; do not translate titles. Connecting words around them ('See', 'cf.', 'For a discussion of ... see') are rendered in Korean (예: '…에 관해서는 다음을 참조: ...').
- Use the glossary (if given) for terms.
- Fix obvious OCR artifacts silently.
"""

NOTES_RULE_COMMENTARY = """Mode: COMMENTARY ONLY. For each note decide:
- If the note is purely bibliographic (only source citations, 'Ibid.', 'See X', 'Op. cit.', page references, with at most a few connecting words), output exactly `[^ID]: DROP`.
- Otherwise (the note contains explanation, argument, data, qualification, or discussion), translate it: commentary in Korean, citations kept in English."""

NOTES_RULE_ALL = """Mode: ALL NOTES. Translate every note (never output DROP): commentary in Korean, citations kept in English."""

NOTES_PER_REQ_CHARS = 9000


def build_notes_system_blocks(glossary: str = "", mode: str = None) -> list:
    glossary = merge_glossary(glossary)
    rule = NOTES_RULE_ALL if (mode or FOOTNOTE_MODE) == "all" else NOTES_RULE_COMMENTARY
    blocks = [{"type": "text", "text": NOTES_SYSTEM.replace("{mode_rule}", rule)}]
    if glossary and glossary.strip():
        blocks.append({"type": "text",
                       "text": "## Glossary (use these renderings consistently)\n" + glossary.strip()})
    blocks[-1]["cache_control"] = {"type": "ephemeral"}
    return blocks


def group_notes(notes: list, max_chars: int = NOTES_PER_REQ_CHARS) -> list:
    """notes: [{'id','num','page','text'}] -> [[note,...], ...] 요청 단위 묶음"""
    groups, cur, n = [], [], 0
    for nt in notes:
        L = len(nt["text"]) + 40
        if cur and n + L > max_chars:
            groups.append(cur); cur, n = [], 0
        cur.append(nt); n += L
    if cur:
        groups.append(cur)
    return groups


def build_notes_user_message(group: list) -> str:
    def esc(t):
        return t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    body = "\n".join(f'<note id="{nt["id"]}" n="{nt["num"]}">{esc(nt["text"])}</note>' for nt in group)
    return "<notes>\n" + body + "\n</notes>"


def parse_notes_output(raw: str, group: list, mode: str = None) -> dict:
    """모델 출력 -> {id: 번역문 또는 None(DROP)}. 누락된 id 는 키 없음."""
    out = {}
    valid = {nt["id"] for nt in group}
    for m in FNDEF_RE.finditer(raw or ""):
        nid, txt = m.group(1), m.group(2).strip()
        if nid not in valid:
            continue
        if re.fullmatch(r"\(?\s*DROP\s*\)?\.?", txt, re.IGNORECASE) or re.search(r"\)\s*DROP\s*$", txt):
            if (mode or FOOTNOTE_MODE) != "all":
                out[nid] = None          # 순수 서지 각주 -> 삭제 (all 모드면 키 없음 -> 원문 유지)
        else:
            out[nid] = txt
    return out


def notes_to_defs(note_map: dict, notes: list) -> str:
    """번역된 각주 -> '[^id]: 텍스트' 정의 블록 (번역 실패분은 원문 영어로)."""
    lines = []
    for nt in notes:
        if nt["id"] in note_map:
            t = note_map[nt["id"]]
            if t:
                lines.append(f"[^{nt['id']}]: {t}")
        else:
            lines.append(f"[^{nt['id']}]: (원주 {nt['num']}) {nt['text']}")
    return "\n".join(lines)


def load_user_glossary() -> str:
    """user_glossary.txt: 'English = 한국어' 줄들 (# 주석 허용)."""
    try:
        if USER_GLOSSARY_PATH is not None and USER_GLOSSARY_PATH.exists():
            out = []
            for ln in USER_GLOSSARY_PATH.read_text(encoding="utf-8-sig").splitlines():
                ln = ln.strip()
                if ln and not ln.startswith("#") and "=" in ln:
                    out.append(ln)
            return "\n".join(out)
    except Exception:
        pass
    return ""


def merge_glossary(auto: str) -> str:
    """사용자 용어집 우선 + 자동 용어집 중 겹치지 않는 항목."""
    user = load_user_glossary()
    if not user:
        return auto or ""
    keys = {ln.split("=", 1)[0].strip().lower() for ln in user.splitlines()}
    rest = [ln for ln in (auto or "").splitlines()
            if "=" in ln and ln.split("=", 1)[0].strip().lower() not in keys]
    return "# user-specified (priority; apply when the term is used in the sense shown)\n" + user + ("\n# auto-extracted\n" + "\n".join(rest) if rest else "")


def build_system_blocks(glossary: str = "") -> list:
    """system 파라미터로 넘길 블록 리스트. 마지막 블록에 cache_control (프롬프트 캐싱)."""
    glossary = merge_glossary(glossary)
    blocks = [{"type": "text", "text": SYSTEM_PROMPT}]
    if glossary and glossary.strip():
        blocks.append({"type": "text",
                       "text": "## Glossary (use these renderings consistently)\n" + glossary.strip()})
    blocks[-1]["cache_control"] = {"type": "ephemeral"}
    return blocks


def build_user_message(chunk: str, idx: int, total: int,
                       prev_src: str = "", prev_tr: str = "") -> str:
    parts = [f"<section index=\"{idx+1}\" of=\"{total}\"/>"]
    if prev_src:
        parts.append("<previous_source>\n" + prev_src.strip() + "\n</previous_source>")
    if prev_tr:
        parts.append("<previous_translation>\n" + prev_tr.strip() + "\n</previous_translation>")
    parts.append("<text_to_translate>\n" + chunk.strip() + "\n</text_to_translate>")
    return "\n\n".join(parts)


def tail(text: str, n: int) -> str:
    """문맥용 꼬리 — 단락 경계에서 자른다."""
    if not text or n <= 0:
        return ""
    t = text.strip()
    if len(t) <= n:
        return t
    cut = t[-n:]
    p = cut.find("\n\n")
    if 0 <= p < n * 0.6:
        cut = cut[p + 2:]
    return cut.strip()


LAST_CITATION = ""


def _parse_glossary(raw: str) -> str:
    """용어집 응답 -> 'English = 한국어' 줄들. 첫 줄 CITATION: 은 LAST_CITATION 으로 따로 보관."""
    global LAST_CITATION
    lines = []
    for ln in (raw or "").splitlines():
        st = ln.strip()
        m = re.match(r"^\W*CITATION\s*:\s*(.+)$", st, re.IGNORECASE)
        if m:
            LAST_CITATION = m.group(1).strip()
            continue
        st = st.lstrip("-*•0123456789. ").strip()
        if "=" in st and len(st) < 200:
            lines.append(st)
    return "\n".join(lines)


def get_citation() -> str:
    return LAST_CITATION


def build_glossary(client, model: str, text: str, max_chars: int = 150_000) -> tuple:
    """용어집 추출 (API 1회). (glossary_text, usage) 반환."""
    src = text if len(text) <= max_chars else (text[:max_chars] + "\n\n[... truncated ...]")
    src = MARKER_RE.sub("", src)
    kwargs = get_thinking_kwargs(model, "low")
    resp = call_with_retry(lambda: client.messages.create(
        model=model,
        max_tokens=6000,
        system=GLOSSARY_SYSTEM,
        messages=[{"role": "user", "content": "<paper>\n" + src + "\n</paper>"}],
        **kwargs,
    ))
    return _parse_glossary(extract_response_text(resp.content)), resp.usage


# ──────────────────────────────────────────────
# 로그 파일 동시 기록 (tee)
# ──────────────────────────────────────────────
class _TeeWriter:
    def __init__(self, stream, log_path):
        self._stream = stream
        try:
            self._log = open(str(log_path), 'w', encoding='utf-8', buffering=1)
        except Exception:
            self._log = None
    def write(self, data):
        try:
            self._stream.write(data)
        except Exception:
            pass
        if self._log:
            try:
                self._log.write(data)
            except Exception:
                pass
    def flush(self):
        try:
            self._stream.flush()
        except Exception:
            pass
        if self._log:
            try:
                self._log.flush()
            except Exception:
                pass
    @property
    def encoding(self):
        return 'utf-8'
    @property
    def errors(self):
        return 'replace'
    def fileno(self):
        return self._stream.fileno()
    def isatty(self):
        try:
            return self._stream.isatty()
        except Exception:
            return False


def setup_tee_log(log_name="last_run.log"):
    sys.stdout = _TeeWriter(sys.stdout, BASE_DIR / log_name)


# ──────────────────────────────────────────────
# 파일 읽기
# ──────────────────────────────────────────────
def assets_dir_for(input_path: str) -> Path:
    stem = re.sub(r"[^\w\-\.\(\)가-힣 ]", "_", Path(input_path).stem)
    return ASSETS_ROOT / stem


def read_document(path: str):
    """(text, assets, stats) — assets: {img_id: {...}} (PDF 외 포맷은 {})."""
    p = Path(path)
    ext = p.suffix.lower()
    if ext == ".pdf":
        from pdf_extract import extract_pdf
        text, assets, stats = extract_pdf(
            str(p), assets_dir_for(path),
            strip_footnotes=STRIP_FOOTNOTES,
            footnote_mode=FOOTNOTE_MODE,
            page_markers=PAGE_MARKERS,
            first_page=FIRST_PAGE,
            extract_figures=EXTRACT_FIGURES,
            extract_tables=EXTRACT_FIGURES,
            extract_equations=EXTRACT_FIGURES,
        )
        return text, assets, stats
    elif ext == ".docx":
        from docx import Document
        doc = Document(str(p))
        text = "\n\n".join(para.text for para in doc.paragraphs if para.text.strip())
        return text, {}, {}
    else:
        return p.read_text(encoding="utf-8", errors="replace"), {}, {}


def read_file(path: str) -> str:
    """하위 호환용."""
    return read_document(path)[0]


# ──────────────────────────────────────────────
# 텍스트 정제 (번역 전)  — pdf_extract 이후의 2차 안전망
# ──────────────────────────────────────────────
def clean_text(text: str) -> str:
    lines = text.split("\n")

    stripped_lines = [l.strip() for l in lines if l.strip()]
    counts = Counter(stripped_lines)
    repeated = {s for s, c in counts.items() if c >= 3 and len(s) <= 120 and not MARKER_RE.match(s)}

    heading_text_counts: Counter = Counter()
    for l in stripped_lines:
        hm = re.match(r'^(#{1,4})\s+(.+)$', l)
        if hm:
            heading_text_counts[hm.group(2).strip()] += 1
    repeated_heading_texts = {t for t, c in heading_text_counts.items() if c >= 2 and len(t) <= 100}

    cleaned = []
    seen_repeated = set()
    for line in lines:
        s = line.strip()
        if not s:
            cleaned.append("")
            continue
        if MARKER_RE.match(s) or PAGE_RE.fullmatch(s):
            cleaned.append(s)
            continue

        hm = re.match(r'^(#{1,4})\s+(.+)$', s)
        if hm and hm.group(2).strip() in repeated_heading_texts:
            continue

        if s in repeated:
            if s not in seen_repeated:
                seen_repeated.add(s)
                cleaned.append(line)
            continue

        s_low = s.lower()
        if any(kw in s_low for kw in [
            "jstor.org", "downloaded from", "all use subject to",
            "content downloaded", "heinonline", "hein.journals",
            "william s. hein", "getzville",
            "imf.org", "imf publication",
            "terms & conditions", "terms and conditions",
        ]):
            continue
        if re.match(r"^(ALWD|APA|MLA|Chicago|McGill Guide|AGLC|OSCOLA)\s+\d", s):
            continue
        if re.match(r"^\d{1,2}/\d{1,2}/\d{2,4}\s+\d{1,2}:\d{2}", s):
            continue
        if re.match(r"^(Sunday|Monday|Tuesday|Wednesday|Thursday|Friday|Saturday),?\s", s):
            continue
        if re.match(r"^\d{1,3}\.\d{1,3}\.\d{1,3}\.\d{1,3}", s):
            continue
        if re.match(r"^https?://", s):
            continue
        if re.match(r"^\d{1,4}$", s):
            continue

        _no_punct = not re.search(r"[.!?;]", s)
        if _no_punct and re.match(r"^\d{1,4}\s+[A-Z][A-Za-z\s\-\'/&]{2,100}$", s):
            continue
        if _no_punct and re.match(r"^[A-Z][A-Za-z\s\-\'/&]{2,100}\s+\d{1,4}$", s):
            continue
        if _no_punct and re.match(
            r"^[A-Z][A-Za-z\s\-\'/&]{2,60}"
            r"\s+(Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)"
            r"[a-z/A-Z\s]*\d{4}\s+\d{1,4}$", s
        ):
            continue
        _months = (r"(January|February|March|April|May|June|July|August|"
                   r"September|October|November|December)")
        if _no_punct and re.match(fr"^{_months}([/\-]{_months})?\s+\d{{4}}$", s, re.IGNORECASE):
            continue
        if re.match(r"^[-*_]{3,}$", s):
            continue
        if re.match(r"^(Hardcover|Paperback)\s+\$[\d.]+", s, re.IGNORECASE):
            continue
        if re.match(r"^Downloaded:?\s*$", s, re.IGNORECASE):
            continue

        cleaned.append(line)

    if STRIP_REFERENCES:
        _REF_PAT = re.compile(
            r"^(#{1,4}\s+)?(\d+\.?\s+)?"
            r"(references|bibliography|works\s+cited|참고\s*문헌)\s*:?\s*$",
            re.IGNORECASE,
        )
        cut_idx = None
        for i, ln in enumerate(cleaned):
            if _REF_PAT.match(ln.strip()):
                cut_idx = i
                break
        if cut_idx is not None:
            cleaned = cleaned[:cut_idx]

    result = "\n".join(cleaned)
    result = re.sub(r"\n{3,}", "\n\n", result)
    return result.strip()


# ──────────────────────────────────────────────
# 청크 분할 (문장 경계 보장)
# ──────────────────────────────────────────────
_SENT_SPLIT = re.compile(r"(?<=[.!?])[\"”’)\]]*\s+(?=[\"“(\[A-Z0-9])")


def _split_long_paragraph(para: str, max_chars: int) -> list:
    """max_chars 를 넘는 단락을 문장 경계에서 나눈다."""
    if len(para) <= max_chars:
        return [para]
    sents = _SENT_SPLIT.split(para)
    out, cur = [], ""
    for s in sents:
        if cur and len(cur) + len(s) + 1 > max_chars:
            out.append(cur.strip())
            cur = s
        else:
            cur = (cur + " " + s) if cur else s
    if cur.strip():
        out.append(cur.strip())
    # 문장 하나가 max_chars 를 넘는 극단적 경우 강제 분할
    final = []
    for piece in out:
        while len(piece) > max_chars * 1.5:
            cut = piece.rfind(" ", 0, max_chars)
            cut = cut if cut > 0 else max_chars
            final.append(piece[:cut].strip())
            piece = piece[cut:].strip()
        final.append(piece)
    return final


def chunk_text(text: str, max_chars: int = MAX_CHUNK) -> list:
    paragraphs = [p.strip() for p in re.split(r"\n{2,}", text) if p.strip()]
    units = []
    for para in paragraphs:
        units.extend(_split_long_paragraph(para, max_chars))

    chunks, current = [], ""
    for u in units:
        if current and len(current) + len(u) + 2 > max_chars:
            chunks.append(current.strip())
            current = u
        else:
            current = (current + "\n\n" + u) if current else u
    if current.strip():
        chunks.append(current.strip())

    # 청크 끝에 헤딩만 남으면 다음 청크 앞으로 이동
    for k in range(len(chunks) - 1):
        parts = chunks[k].rsplit("\n\n", 1)
        if len(parts) == 2 and parts[1].lstrip().startswith("#"):
            chunks[k] = parts[0].strip()
            chunks[k + 1] = parts[1].strip() + "\n\n" + chunks[k + 1]

    # 너무 짧은 꼬리는 앞 청크에 합침 (단, 마커만 있는 청크는 그대로 둔다)
    out = []
    for c in chunks:
        if out and len(c) < MIN_CHUNK and not MARKER_RE.search(c) and not PAGE_RE.search(c):
            out[-1] = out[-1] + "\n\n" + c
        else:
            out.append(c)
    return out


def chunk_needs_api(chunk: str) -> bool:
    """마커/공백만 있는 청크는 API 호출 없이 그대로 통과."""
    rest = PAGE_RE.sub("", MARKER_RE.sub("", chunk))
    rest = FNREF_RE.sub("", rest).strip()
    return len(rest) >= 3


# ──────────────────────────────────────────────
# 번역 결과 후처리
# ──────────────────────────────────────────────
def clean_translated(text: str) -> str:
    stripped = text.strip()
    if stripped in ("(empty string)", "empty string", "(빈 문자열)", ""):
        return ""
    reject_phrases = [
        "번역할 본문 내용이 제공되지 않았습니다",
        "번역하고자 하시는 학술 텍스트를",
        "한국어로 번역해 드리겠습니다",
        "주어진 텍스트에는 번역할",
        "please provide",
        "it seems the text",
        "no content to translate",
        "(empty string)",
    ]
    if any(p in text for p in reject_phrases) and len(text) < 400:
        return ""

    # 모델이 태그를 되돌려준 경우 제거
    text = re.sub(r"</?(text_to_translate|previous_source|previous_translation|section)[^>]*>", "", text)

    meta_line_patterns = [
        re.compile(r"^\s*Human:\s*", re.IGNORECASE),
        re.compile(r"^\s*Assistant:\s*", re.IGNORECASE),
        re.compile(r"Per Rule\s+\d+", re.IGNORECASE),
        re.compile(r"\bskip(?:ped)?\s+entirely\b", re.IGNORECASE),
        re.compile(r"The input is\s+(citation|reference|database|metadata|boilerplate)", re.IGNORECASE),
        re.compile(r"^\s*\(?\s*per rule\b", re.IGNORECASE),
        re.compile(r"이 (내용|입력|자료|텍스트)은.*?(건너뜁|스킵|skip)", re.IGNORECASE),
    ]
    cleaned_lines = []
    for ln in text.splitlines():
        if any(p.search(ln) for p in meta_line_patterns):
            continue
        cleaned_lines.append(ln)
    out = "\n".join(cleaned_lines)

    # 마커 정규화: [[ IMG : id ]] -> [[IMG:id]], 항상 독립 줄
    out = MARKER_RE.sub(lambda m: "\n\n[[IMG:" + m.group(1) + "]]\n\n", out)
    # 쪽 마커 정규화 (모델이 [원서 p. N] 로 바꿔 쓴 경우도 복구)
    out = re.sub(r"\[\s*원서\s*p\.?\s*(\d{1,4})\s*\]", r"[[PAGE:\1]]", out)
    out = PAGE_RE.sub(lambda m: "[[PAGE:" + m.group(1) + "]]", out)
    out = re.sub(r"`(\[\[PAGE:[^\]]+\]\])`", r"\1", out)
    # 각주 참조 앞 공백 제거
    out = re.sub(r"[ \t]+(\[\^[A-Za-z0-9_\-]+\])(?!:)", r"\1", out)

    # 각주 콜아웃 잔재 정리 (보수적 패턴)
    out = re.sub(r'(["”’\)\]])\s*\d{1,3}(?=\s+[가-힣A-Z])', r'\1', out)
    out = re.sub(r'([.!?])\s+\d{1,3}(?=\s+[가-힣A-Z])', r'\1', out)
    out = re.sub(r'([가-힣A-Za-z])\.\d{1,3}(?=\s+[가-힣A-Z])', r'\1.', out)

    out = re.sub(r"\n{3,}", "\n\n", out).strip()
    return out


def _insert_at_fraction(text: str, token: str, frac: float) -> str:
    """번역문의 비슷한 상대 위치(문장 끝 우선)에 토큰을 끼워 넣는다."""
    if not text:
        return token
    target = int(len(text) * max(0.0, min(1.0, frac)))
    best = None
    for m in re.finditer(r"(?<=[.!?。다])[\"”’)\]]*(?=\s)", text):
        d = abs(m.end() - target)
        if best is None or d < best[0]:
            best = (d, m.end())
    pos = best[1] if best else len(text)
    sep = "" if token.startswith("[^") else " "
    return text[:pos] + sep + token + text[pos:]


def restore_markers(translated: str, source_chunk: str) -> tuple:
    """원문 청크의 이미지/쪽/각주 토큰이 번역문에서 빠졌으면 복원. (text, missing_ids) 반환.
    이미지 마커는 청크 끝에, 쪽 마커·각주 참조는 원문상의 상대 위치에 맞춰 넣는다."""
    missing = []
    expected = MARKER_RE.findall(source_chunk)
    if expected:
        present = set(MARKER_RE.findall(translated))
        miss = [m for m in expected if m not in present]
        if miss:
            translated = translated.rstrip() + "\n\n" + "\n\n".join(f"[[IMG:{m}]]" for m in miss)
            missing += miss
    L = max(1, len(source_chunk))
    present_pages = set(PAGE_RE.findall(translated))
    for m in PAGE_RE.finditer(source_chunk):
        if m.group(1) not in present_pages:
            tok = f"[[PAGE:{m.group(1)}]]"
            before = source_chunk[max(0, m.start() - 2):m.start()]
            if before.endswith("\n\n") or m.start() == 0:
                # 독립 줄 마커: 같은 상대 위치의 단락 경계에 둔다
                paras = translated.split("\n\n")
                k = min(len(paras), round(len(paras) * m.start() / L))
                paras.insert(k, tok)
                translated = "\n\n".join(paras)
            else:
                translated = _insert_at_fraction(translated, tok, m.start() / L)
            missing.append(tok)
    present_refs = set(FNREF_RE.findall(translated))
    for m in FNREF_RE.finditer(source_chunk):
        if m.group(1) not in present_refs:
            tok = f"[^{m.group(1)}]"
            translated = _insert_at_fraction(translated, tok, m.end() / L)
            missing.append(tok)
    return translated, missing


def _is_issue_date(text: str) -> bool:
    if re.match(r"^\d{4}년\s+\d{1,2}([/\-]\d{1,2})?월호?$", text.strip()):
        return True
    _m = (r"(January|February|March|April|May|June|July|August|"
          r"September|October|November|December)")
    if re.match(fr"^{_m}([/\-]{_m})?\s+\d{{4}}$", text.strip(), re.IGNORECASE):
        return True
    return False


# ──────────────────────────────────────────────
# 가격표 / 비용 추적  (USD per MTok, platform.claude.com/docs/en/about-claude/pricing 2026-10)
# ──────────────────────────────────────────────
PRICE = [
    # (substring, input, output, cache_write_5m, cache_read)  — 위에서부터 먼저 매칭
    ("fable-5-1",   10.0, 50.0, 12.50, 0.25),
    ("mythos-5-1",  10.0, 50.0, 12.50, 0.25),
    ("fable",       10.0, 50.0, 12.50, 1.00),
    ("mythos",      10.0, 50.0, 12.50, 1.00),
    ("opus-5-5",     4.0, 20.0,  5.00, 0.20),
    ("opus-5",       5.0, 25.0,  6.25, 0.50),
    ("opus-4-8",     5.0, 25.0,  6.25, 0.50),
    ("opus-4-7",     5.0, 25.0,  6.25, 0.50),
    ("opus-4-6",     5.0, 25.0,  6.25, 0.50),
    ("opus",        15.0, 75.0, 18.75, 1.50),   # opus-4-5 이하
    ("sonnet-5",     2.0, 10.0,  2.50, 0.20),   # sonnet-5 / sonnet-5-5
    ("sonnet",       3.0, 15.0,  3.75, 0.30),   # sonnet-4-x
    ("haiku-4-5",    1.0,  5.0,  1.25, 0.10),
    ("haiku",        0.8,  4.0,  1.00, 0.08),
]


def _model_prices(model: str):
    m = (model or "").lower()
    for key, i, o, cw, cr in PRICE:
        if key in m:
            return i, o, cw, cr
    return 3.0, 15.0, 3.75, 0.30


class CostTracker:
    def __init__(self, model: str, batch_discount: bool = False):
        self.model    = model
        self.discount = 0.5 if batch_discount else 1.0
        self.input_tokens  = 0
        self.output_tokens = 0
        self.cache_write   = 0
        self.cache_read    = 0
        self.calls = 0

    def add(self, usage):
        self.input_tokens  += getattr(usage, "input_tokens", 0) or 0
        self.output_tokens += getattr(usage, "output_tokens", 0) or 0
        self.cache_write   += getattr(usage, "cache_creation_input_tokens", 0) or 0
        self.cache_read    += getattr(usage, "cache_read_input_tokens", 0) or 0
        self.calls += 1

    def total_cost(self):
        ip, op, cw, cr = _model_prices(self.model)
        d = self.discount
        ic = self.input_tokens / 1e6 * ip * d
        oc = self.output_tokens / 1e6 * op * d
        cc = (self.cache_write / 1e6 * cw + self.cache_read / 1e6 * cr) * d
        return ic, oc, cc

    def report(self):
        ic, oc, cc = self.total_cost()
        label = " (batch 50% off)" if self.discount < 1 else ""
        lines = [
            "",
            "=" * 50,
            "  Cost Report" + label,
            "=" * 50,
            "  Model:         " + self.model,
            "  API calls:     " + str(self.calls),
            "  Input tokens:  " + f"{self.input_tokens:,}",
            "  Input cost:    $" + f"{ic:.4f}",
            "  Output tokens: " + f"{self.output_tokens:,}",
            "  Output cost:   $" + f"{oc:.4f}",
        ]
        if self.cache_write or self.cache_read:
            lines += [
                "  Cache write:   " + f"{self.cache_write:,}" + "  read: " + f"{self.cache_read:,}",
                "  Cache cost:    $" + f"{cc:.4f}",
            ]
        lines += [
            "  Total:         $" + f"{ic + oc + cc:.4f}",
            "=" * 50,
        ]
        for line in lines:
            print(line, flush=True)
        try:
            (BASE_DIR / "last_cost.txt").write_text("\n".join(lines), encoding="utf-8")
        except Exception:
            pass


# ──────────────────────────────────────────────
# 번역문 내 러닝 헤더 잔재 제거 (저장 직전 최종 필터)
# ──────────────────────────────────────────────
def _fuzzy_name_pattern(tokens):
    """이름 토큰 시퀀스를 구분자(공백/·/,/과/와 등) 허용 정규식으로."""
    sep = r"[\s\n·ㆍ,\.&*()\-]{0,4}(?:와|과|및|and)?[\s\n·ㆍ,\.&*()\-]{0,4}"
    return sep.join(re.escape(t) for t in tokens)


def strip_header_leakage(full: str) -> str:
    """모델이 놓친 러닝 헤더(페이지번호+저자명/축약 제목)를 번역문에서 제거.

    1) '518 김명섭·호레이스 제프리 하지스' 같은 독립 헤더 줄 제거
    2) 문서 첫머리의 저자 표시줄에서 저자 이름 시퀀스를 추출해,
       본문 문장 이음새에 끼어든 저자명 나열(첫 등장 제외)을 제거
    """
    lines = full.split("\n")

    # ── 1) 독립 헤더 줄: 2-4자리 페이지 번호 + 짧은 텍스트, 문장부호 없음 ──
    kept = []
    for l in lines:
        t = l.strip()
        drop = False
        if t and not t.startswith("#") and len(t) <= 70 \
                and not re.search(r"[.!?;]", t) and not MARKER_RE.search(t) and not PAGE_RE.search(t):
            if re.match(r"^\d{2,4}\s+\S", t) or re.search(r"\s\d{2,4}$", t):
                core = re.sub(r"^\d{2,4}\s+|\s+\d{2,4}$", "", t)
                if core and len(core.split()) <= 10 and not re.match(r"^[\d\W]+$", core):
                    drop = True
        if not drop:
            kept.append(l)
    full = "\n".join(kept)

    # ── 2) 저자 표시줄 추출: 앞부분에서 헤딩 직후에 오는 '이름꼴 줄'을 수집 ──
    #    (표지/목차가 앞에 붙은 책 챕터 스캔도 지원: 첫 헤딩만 보지 않고 첫 60줄을 훑는다)
    author_lines = []
    byline_end = 0
    offset = 0
    lines_since_heading = 99
    for l in full.split("\n")[:60]:
        line_end = offset + len(l) + 1
        t = l.strip()
        if not t:
            offset = line_end
            continue
        is_heading = t.startswith("#")
        body = re.sub(r"^#+\s*", "", t)
        core = re.sub(r"\([^)]*\)", " ", body)        # (WILLIAMSON MURRAY) 등 괄호 병기 제거
        ntok = len(re.findall(r"[가-힣A-Za-z\.]{2,}", core))
        name_like = (len(body) <= 90 and not re.search(r"[.!?:;\d]", core)
                     and re.match(r"^[가-힣A-Za-z][가-힣A-Za-z\s·ㆍ,\.\*\-]*$", core.strip() or "-"))
        if is_heading:
            if name_like and ntok >= 2 and lines_since_heading <= 2:
                author_lines.append(core)              # 헤딩으로 표기된 저자 표시줄
                byline_end = line_end
            lines_since_heading = 0
        elif name_like and ntok >= 1 and lines_since_heading <= 3:
            author_lines.append(core)
            byline_end = line_end
        else:
            lines_since_heading += 1
        offset = line_end

    # 저자 토큰: 연결어(와/과/및/and/,/·)로 저자 단위 분리 → 전체 시퀀스 + 성(姓) 시퀀스
    all_tokens = []
    surnames = []
    for line in author_lines:
        for author in re.split(r"\s*(?:,|·|ㆍ|\band\b)\s*|\s+(?:와|과|및)\s+|(?<=[가-힣])(?:와|과)(?=\s)", line):
            toks = [w.rstrip("*") for w in re.findall(r"[가-힣A-Za-z\.]{2,}", author)]
            toks = [w for w in toks if w]
            if toks:
                all_tokens.extend(toks)
                surnames.append(toks[-1])   # 한국 학술지 관행상 헤더는 성(마지막 토큰) 나열

    patterns = []
    if len(all_tokens) >= 2 and sum(len(w) for w in all_tokens) >= 6:
        patterns.append(_fuzzy_name_pattern(all_tokens))
    if len(surnames) >= 2 and surnames != all_tokens:
        patterns.append(_fuzzy_name_pattern(surnames))

    for p in patterns:
        pat = re.compile(r"(?:\d{1,4}\s{0,2})?" + p + r"(?:\s{0,2}\d{1,4})?")
        out_parts = []
        pos = 0
        for m in pat.finditer(full):
            before = full[max(0, m.start() - 3):m.start()]
            after = full[m.end():m.end() + 2]
            if m.start() < byline_end:
                continue                     # 문서 첫머리 저자 표시줄 보존
            seg = m.group(0)
            has_num = bool(re.search(r"\d", seg))
            at_seam = bool(re.search(r"[.!?\"”’)\]]\s*$", before)) or before.endswith("\n") or m.start() == 0
            line_alone = (before.endswith("\n") or m.start() == 0) and (after.startswith("\n") or after == "")
            if has_num or at_seam or line_alone:
                out_parts.append(full[pos:m.start()])
                pos = m.end()
        out_parts.append(full[pos:])
        full = "".join(out_parts)

    # ── 3) 해석 불가능한 잔여 대괄호 마커 제거 ([[IMG:id]] 정상 마커는 보존) ──
    def _drop_bad_marker(m):
        return m.group(0) if (MARKER_RE.fullmatch(m.group(0)) or PAGE_RE.fullmatch(m.group(0))) else ""
    full = re.sub(r"\[\[[^\[\]\n]{0,60}\]\]", _drop_bad_marker, full)

    full = re.sub(r"[ \t]{2,}", " ", full)
    full = re.sub(r"[ \t]+\n", "\n", full)
    full = re.sub(r"\n{3,}", "\n\n", full)
    return full


# ──────────────────────────────────────────────
# 저장 공통: 번역 청크 -> 블록 스트림
# ──────────────────────────────────────────────
def _bind_vision_endnotes(full: str) -> str:
    """비전 출력의 미주 [^eN] 참조/정의를 위치 순서로 짝지어 고유 id 로 바꾼다.
    정의 줄은 같은 번호의 '가장 최근 미연결 참조'와 연결된다(장마다 번호가 다시 시작해도 안전)."""
    tok = re.compile(r"^[ \t]*\[\^v(\d{1,3})\]:|\[\^v(\d{1,3})\](?!:)", re.MULTILINE)
    pending = {}            # num -> [ref match spans...]
    repl = {}               # span -> new text
    k = 0
    for m in tok.finditer(full):
        if m.group(2) is not None:
            pending.setdefault(m.group(2), []).append(m.span())
        else:
            num = m.group(1)
            if pending.get(num):
                k += 1
                ref_span = pending[num].pop()
                pending[num] = []           # 같은 번호의 더 이전 참조는 다른 장/쪽 것으로 보고 폐기
                repl[ref_span] = f"[^b{k}]"
                repl[m.span()] = f"[^b{k}]:"
            else:
                repl[m.span()] = "[^orphan]:"
    if not repl and "[^v" not in full:
        return full
    out, pos = [], 0
    for m in tok.finditer(full):
        out.append(full[pos:m.start()])
        out.append(repl.get(m.span(), ""))        # 짝 없는 참조는 제거
        pos = m.end()
    out.append(full[pos:])
    return "".join(out)


def assemble_full(translated_chunks: list) -> tuple:
    """번역 청크들 -> (본문 전체 텍스트, 각주 {id: 번역문}). 각주 정의 줄은 본문에서 분리."""
    cleaned = [clean_translated(c) for c in translated_chunks]
    full = "\n\n".join(c for c in cleaned if c.strip())
    full = _bind_vision_endnotes(full)
    notes = {}
    def _grab(m):
        t = m.group(2).strip()
        if t and m.group(1) != "orphan":
            notes[m.group(1)] = t
        return ""
    full = FNDEF_RE.sub(_grab, full)
    full = strip_header_leakage(full)
    # 정의 없는 각주 참조 제거
    full = FNREF_RE.sub(lambda m: m.group(0) if m.group(1) in notes else "", full)
    # 중복 쪽 마커(같은 번호 연속) 정리: 첫 번째만 남김
    seen = set()
    def _dedup(m):
        n = m.group(1)
        if n in seen:
            return ""
        seen.add(n)
        return m.group(0)
    full = PAGE_RE.sub(_dedup, full)
    # 빈 헤딩 줄('##' 만 남은 것) 제거
    full = re.sub(r"(?m)^#{1,6}[ \t]*$", "", full)
    full = re.sub(r"[ \t]{2,}", " ", full)
    full = re.sub(r"\n{3,}", "\n\n", full).strip()
    return full, notes


_CAPTION_KO_RE = re.compile(r"^\**\s*(그림|표|도표|지도|사진|Figure|Table|Fig\.)\s*[\dIVX]+[\.\-\d]*[a-z]?[\.:)]?\s", re.IGNORECASE)


def _iter_blocks(translated_chunks=None, full: str = None):
    """(kind, payload) 시퀀스: title / author / heading(level, text) / list / quote / image(id) / caption /
    table / page(N) / para(text).  텍스트에는 인라인 [[PAGE:N]] / [^id] 가 남아 있다."""
    if full is None:
        full, _ = assemble_full(translated_chunks)
    state = {"after_title": False, "seen_title": False}

    def _split_edges(block):
        pre, post = [], []
        while True:
            m = re.match(r"^\s*\[\[PAGE:([^\]]+)\]\]\s*", block)
            if not m:
                break
            pre.append(m.group(1)); block = block[m.end():]
        while True:
            m = re.search(r"\s*\[\[PAGE:([^\]]+)\]\]\s*$", block)
            if not m or m.start() == 0:
                break
            post.insert(0, m.group(1)); block = block[:m.start()]
        return pre, block.strip(), post

    prev_kind = None
    last_heading = None
    held = []          # 문서 첫머리 쪽 마커는 제목·저자 줄 뒤로 미룬다 (샘플 형식)
    started = False
    for raw in re.split(r"\n{2,}", full):
        pre, block, post = _split_edges(raw.strip())
        for n in pre:
            if not started:
                held.append(n)
            else:
                yield ("page", n)
        if block:
            for kind, payload in _block_kinds(block, state):
                if kind == "heading" and prev_kind == "heading" and last_heading == payload[1].casefold():
                    continue                       # 같은 제목이 연달아 두 번 나온 경우
                if kind == "heading":
                    last_heading = payload[1].casefold()
                if not started and kind not in ("title", "author"):
                    started = True
                    for n in held:
                        yield ("page", n)
                    held = []
                # 그림 바로 앞/뒤의 '그림 N. ...' 단락은 캡션 스타일
                if kind == "para" and _CAPTION_KO_RE.match(PAGE_RE.sub("", payload).strip()) and len(payload) <= 300:
                    kind = "caption"
                yield (kind, payload)
                prev_kind = kind
        for n in post:
            yield ("page", n)


def _block_kinds(block, state):
    lines = [l.strip() for l in block.split("\n") if l.strip()]
    if not lines:
        return
    first_line = lines[0]

    mk = MARKER_RE.fullmatch(first_line)
    if mk and len(lines) == 1:
        yield ("image", mk.group(1))
        return
    if MARKER_RE.search(block):
        pos = 0
        for m in MARKER_RE.finditer(block):
            pre = block[pos:m.start()].strip()
            if pre:
                yield ("para", " ".join(l.strip() for l in pre.split("\n") if l.strip()))
            yield ("image", m.group(1))
            pos = m.end()
        rest = block[pos:].strip()
        if rest:
            yield ("para", " ".join(l.strip() for l in rest.split("\n") if l.strip()))
        return

    if re.match(r"^[-*_]{3,}$", first_line):
        return

    # GFM 표: 표 그대로 단락 텍스트로 (렌더러에서 처리)
    if first_line.startswith("|") and len(lines) >= 2 and re.match(r"^\|?\s*:?-{3,}", lines[1]):
        yield ("table", lines)
        return

    hm = re.match(r"^(#{1,4})\s+(.+)$", first_line)
    if hm:
        level = len(hm.group(1))
        txt = hm.group(2).strip()
        if _is_issue_date(txt):
            return
        plain = PAGE_RE.sub("", txt).strip()
        if not plain:
            return
        if re.fullmatch(r"(주|주석|미주|원주|각주|notes?|endnotes)", plain, re.IGNORECASE):
            return                      # 장 끝 미주 절 제목 — 내용은 각주로 옮겨졌거나 생략됨
        if re.match(r"^\d{1,4}\s+\S", plain) and not re.search(r"[.!?;,]", plain) \
                and len(plain) <= 70 and len(plain.split()) <= 6 and not re.match(r"^\d+(\.\d+)*\s+[가-힣]", plain):
            return
        if level == 1 or not state["seen_title"]:
            state["seen_title"] = True
            state["after_title"] = True
            yield ("title", txt)
        else:
            state["after_title"] = False
            yield ("heading", (level, txt))
        if len(lines) > 1:
            yield ("para", " ".join(lines[1:]))
        return

    lm = re.match(r"^[-*•]\s+(.+)$", first_line)
    if lm:
        items, cur = [], None
        for l in lines:
            m = re.match(r"^[-*•]\s+(.+)$", l)
            if m:
                if cur: items.append(cur)
                cur = m.group(1)
            else:
                cur = (cur + " " + l) if cur else l
        if cur: items.append(cur)
        for it in items:
            yield ("list", it)
        return

    qm = re.match(r"^>\s*(.+)$", first_line)
    if qm:
        yield ("quote", " ".join(l.lstrip(">").strip() for l in lines if l.strip()))
        return

    merged = " ".join(lines)
    # 안전망: 모델이 '##' 를 빠뜨린 절 제목 (짧고, 문장부호로 끝나지 않는 한 줄)
    plain = PAGE_RE.sub("", merged).strip()
    if (not state["after_title"] and state["seen_title"] and len(lines) == 1 and 2 <= len(plain) <= 45
            and not re.search(r"[.,;:!?。…다요\]\"”’]$", plain) and not FNREF_RE.search(plain)
            and not re.match(r"^[\(\[\d①-⑳•\-–—'\"‘“]", plain) and not _CAPTION_KO_RE.match(plain)
            and "," not in plain and len(plain.split()) <= 6):
        yield ("heading", (2, merged))
        return
    if state["after_title"] and len(merged) <= 120 and not FNREF_RE.search(merged) \
            and not re.search(r"[.!?。]\s*$", merged) and not re.search(r"다\.?$", merged):
        state["after_title"] = False
        yield ("author", merged)
        return
    state["after_title"] = False
    yield ("para", merged)


def _image_size_cm(asset: dict, max_w_cm: float = 15.0):
    """원본 pt 크기 기준 삽입 크기(cm). 텍스트 폭 초과 시 축소."""
    w_pt = float(asset.get("width_pt", 300)) or 300.0
    h_pt = float(asset.get("height_pt", 200)) or 200.0
    kind = asset.get("kind", "fig")
    scale = 1.0 if kind == "eq" else 1.0
    w_cm = w_pt / 72.0 * 2.54 * scale
    h_cm = h_pt / 72.0 * 2.54 * scale
    if w_cm > max_w_cm:
        h_cm = h_cm * (max_w_cm / w_cm)
        w_cm = max_w_cm
    return w_cm, h_cm


# ──────────────────────────────────────────────
# 저장 (DOCX / PDF / MD)
# ──────────────────────────────────────────────
def source_line_text(citation: str = "") -> str:
    """문서 머리의 회색 서지 줄: '<서지> — 한국어 번역(본문 전문, 각주 방식, 그림은 원서 이미지)'."""
    fn = {"strip": "각주 생략", "commentary": "논평성 각주만 번역", "all": "각주 전체 번역"}.get(FOOTNOTE_MODE, "")
    tail_ = "한국어 번역(본문 전문" + (", " + fn if fn else "") + (", 그림은 원서 이미지" if EXTRACT_FIGURES else "") + ")"
    citation = (citation or "").strip()
    return (citation + " — " + tail_) if citation else tail_


def save_docx(translated_chunks: list, output_path: str, assets: dict = None, full=None, notes=None,
              citation: str = "") -> int:
    from doc_render import build_docx
    if full is None:
        full, notes = assemble_full(translated_chunks)
    return build_docx(_iter_blocks(full=full), notes or {}, output_path, assets or {}, _image_size_cm,
                      source_line=source_line_text(citation))


def save_md(translated_chunks: list, output_path: str, assets: dict = None, full=None, notes=None,
            citation: str = "") -> None:
    assets = assets or {}
    if full is None:
        full, notes = assemble_full(translated_chunks)
    notes = notes or {}
    def _rep(m):
        a = assets.get(m.group(1))
        if a:
            return f"![{m.group(1)}]({Path(a['path']).as_posix()})"
        return m.group(0)
    body = MARKER_RE.sub(_rep, full)
    body = "> " + source_line_text(citation) + "\n\n" + body
    body = PAGE_RE.sub(lambda m: f"`[원서 p. {m.group(1)}]`", body)
    used = FNREF_RE.findall(body)
    defs = [f"[^{i}]: {notes[i]}" for i in dict.fromkeys(used) if i in notes]
    if defs:
        body += "\n\n" + "\n\n".join(defs)
    Path(output_path).write_text(body + "\n", encoding="utf-8")


def _md_to_rl(text: str) -> str:
    text = re.sub(r"<sup>\s*[\d,\s\-]+\s*</sup>", "", text, flags=re.IGNORECASE)
    text = re.sub(r"[⁰-⁹²³¹]+", "", text)
    text = (text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"))
    text = re.sub(r"\*\*\*(.+?)\*\*\*", r"<b><i>\1</i></b>", text)
    text = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", text)
    text = re.sub(r"\*(.+?)\*", r"<i>\1</i>", text)
    return text


def save_pdf_reportlab(translated_chunks: list, output_path: str, assets: dict = None,
                       full=None, notes=None, citation: str = "") -> None:
    """Word/LibreOffice 가 없을 때의 폴백. 각주는 문서 끝 '원주' 절로 모은다."""
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.units import cm
        from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT, TA_JUSTIFY
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Image, Spacer, Table, TableStyle
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.lib import colors
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
    except ImportError:
        raise ImportError(
            "PDF 출력에는 'reportlab' 패키지가 필요합니다.\n"
            "  설치: pip install reportlab\n"
            "  대안: 출력 형식을 DOCX 로 (bat: Options > 7 Output format / GUI: 왼쪽 [출력 형식])"
        )
    assets = assets or {}
    if full is None:
        full, notes = assemble_full(translated_chunks)
    notes = notes or {}

    font_dir = os.path.join(os.environ.get("SystemRoot", r"C:\Windows"), "Fonts")
    malgun      = os.path.join(font_dir, "malgun.ttf")
    malgun_bold = os.path.join(font_dir, "malgunbd.ttf")
    if not os.path.exists(malgun):
        for cand in ["/usr/share/fonts/truetype/nanum/NanumGothic.ttf",
                     "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
                     "/usr/share/fonts/truetype/noto/NotoSansCJK-Regular.ttc"]:
            if os.path.exists(cand):
                malgun = cand; break
    if os.path.exists(malgun) and "Malgun" not in pdfmetrics.getRegisteredFontNames():
        try:
            pdfmetrics.registerFont(TTFont("Malgun", malgun))
        except Exception:
            pass
    if os.path.exists(malgun_bold) and "Malgun-Bold" not in pdfmetrics.getRegisteredFontNames():
        try:
            pdfmetrics.registerFont(TTFont("Malgun-Bold", malgun_bold))
        except Exception:
            pass
    reg = pdfmetrics.getRegisteredFontNames()
    font_normal = "Malgun" if "Malgun" in reg else "Helvetica"
    font_bold   = "Malgun-Bold" if "Malgun-Bold" in reg else (font_normal if font_normal == "Malgun" else "Helvetica-Bold")

    BODY = 11
    LEAD = BODY * 1.6
    style_title = ParagraphStyle("title", fontName=font_bold, fontSize=18,
        alignment=TA_CENTER, spaceBefore=0, spaceAfter=8, leading=25)
    style_subtitle = ParagraphStyle("subtitle", fontName=font_bold, fontSize=12,
        alignment=TA_CENTER, spaceBefore=2, spaceAfter=6, leading=17)
    style_author = ParagraphStyle("author", fontName=font_bold, fontSize=10,
        alignment=TA_CENTER, spaceAfter=10, leading=15)
    style_caption = ParagraphStyle("caption", fontName=font_bold, fontSize=10,
        alignment=TA_CENTER, spaceAfter=14, leading=15, wordWrap="CJK")
    style_h1 = ParagraphStyle("h1", fontName=font_bold, fontSize=14,
        alignment=TA_LEFT, spaceBefore=16, spaceAfter=8, leading=19)
    style_h2 = ParagraphStyle("h2", fontName=font_bold, fontSize=12,
        alignment=TA_LEFT, spaceBefore=12, spaceAfter=6, leading=17)
    style_h3 = ParagraphStyle("h3", fontName=font_bold, fontSize=11,
        alignment=TA_LEFT, spaceBefore=10, spaceAfter=4, leading=16)
    style_body = ParagraphStyle("body", fontName=font_normal, fontSize=BODY,
        alignment=TA_JUSTIFY, spaceAfter=7, leading=LEAD, wordWrap="CJK")
    style_quote = ParagraphStyle("quote", parent=style_body, leftIndent=cm, rightIndent=cm,
        fontSize=10.5, leading=10.5*1.4, spaceBefore=5, spaceAfter=5)
    style_list = ParagraphStyle("list", parent=style_body, leftIndent=0.6*cm, spaceAfter=3)
    style_marker = ParagraphStyle("marker", fontName=font_normal, fontSize=8.5, textColor=colors.grey,
        alignment=TA_RIGHT, spaceBefore=2, spaceAfter=2, leading=11, wordWrap="CJK")
    style_note = ParagraphStyle("note", fontName=font_normal, fontSize=9, alignment=TA_JUSTIFY,
        leading=9*1.4, spaceAfter=3, wordWrap="CJK")

    doc = SimpleDocTemplate(output_path, pagesize=A4,
        topMargin=2.5*cm, bottomMargin=2.5*cm, leftMargin=2.5*cm, rightMargin=2.5*cm,
        title="Translated Paper")
    frame_w = A4[0] - 5.0*cm

    def _draw_footer(canvas, doc_):
        canvas.saveState()
        canvas.setFont(font_normal, 9)
        canvas.setFillGray(0.4)
        canvas.drawCentredString(A4[0] / 2.0, 1.2 * cm, "- " + str(doc_.page) + " -")
        canvas.restoreState()

    note_order = []
    def _rich(t):
        t = _md_to_rl(t)
        t = PAGE_RE.sub(lambda m: f' <font size="8.5" color="#808080">[원서 p. {m.group(1)}]</font> ', t)
        def _ref(m):
            nid = m.group(1)
            if nid not in notes:
                return ""
            if nid not in note_order:
                note_order.append(nid)
            return f"<super>{note_order.index(nid) + 1}</super>"
        return FNREF_RE.sub(_ref, t)

    story = []
    src_line = source_line_text(citation)
    src_done = False
    for kind, payload in _iter_blocks(full=full):
        if not src_done and kind not in ("title", "author"):
            story.append(Paragraph(_md_to_rl(src_line), style_marker)); src_done = True
        if kind == "caption":
            story.append(Paragraph(_rich(payload), style_caption)); continue
        if kind == "page":
            story.append(Paragraph(f"[원서 p. {payload}]", style_marker)); continue
        if kind == "image":
            a = assets.get(payload)
            if a and Path(a["path"]).exists():
                w_cm, h_cm = _image_size_cm(a, max_w_cm=frame_w / cm * 0.92)
                try:
                    img = Image(a["path"], width=w_cm*cm, height=h_cm*cm)
                    img.hAlign = "CENTER"
                    story.append(Spacer(1, 4)); story.append(img); story.append(Spacer(1, 6))
                except Exception as e:
                    story.append(Paragraph(_md_to_rl(f"[이미지 삽입 실패: {payload}: {e}]"), style_body))
            else:
                story.append(Paragraph(_md_to_rl(f"[이미지 누락: {payload}]"), style_body))
            continue
        if kind == "table":
            rows = [[Paragraph(_rich(c.strip()), style_note) for c in r.strip().strip("|").split("|")]
                    for k, r in enumerate(payload) if k != 1]
            ncol = max(len(r) for r in rows)
            rows = [r + [""] * (ncol - len(r)) for r in rows]
            t = Table(rows, colWidths=[frame_w / ncol] * ncol)
            t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
                                   ("VALIGN", (0, 0), (-1, -1), "TOP")]))
            story.append(t); story.append(Spacer(1, 6)); continue
        if kind == "title":
            story.append(Paragraph(_rich(payload), style_title)); continue
        if kind == "subtitle":
            story.append(Paragraph(_rich(payload), style_subtitle)); continue
        if kind == "heading":
            level, txt = payload
            st = style_h1 if level <= 2 else (style_h2 if level == 3 else style_h3)
            story.append(Paragraph(_rich(txt), st)); continue
        if kind == "list":
            story.append(Paragraph("\u2022 " + _rich(payload), style_list)); continue
        if kind == "quote":
            story.append(Paragraph(_rich(payload), style_quote)); continue
        if kind == "author":
            story.append(Paragraph(_rich(payload), style_author)); continue
        story.append(Paragraph(_rich(payload), style_body))

    if note_order:
        story.append(Spacer(1, 12))
        story.append(Paragraph("원주", style_h2))
        for k, nid in enumerate(note_order, start=1):
            story.append(Paragraph(f"{k}. " + _md_to_rl(notes[nid]), style_note))

    doc.build(story, onFirstPage=_draw_footer, onLaterPages=_draw_footer)


def output_paths(input_file: str, fmt: str, page_range: tuple = None) -> dict:
    """{'docx': path, 'pdf': path, 'md': path} 중 fmt 에 해당하는 것.
    파일명: <원본 stem>[_pp<a>-<b>]_한국어번역.<ext>"""
    p = Path(input_file)
    stem = p.stem
    rng = ""
    if page_range and all(page_range) and not re.search(r"pp?\.?\s*\d", stem):
        rng = f"_pp{page_range[0]}-{page_range[1]}"
    base = p.parent / f"{stem}{rng}_한국어번역"
    fmts = {"docx+pdf": ["docx", "pdf"]}.get(fmt, [fmt])
    return {f: str(base) + "." + f for f in fmts}


def _writable_path(path: str) -> str:
    """대상 파일이 다른 프로그램(PDF 뷰어·Word)에 열려 잠겨 있으면 새 이름으로 바꾼다.
    (2026-10-10 실패 사례: 이전 결과 PDF 가 뷰어에 열려 있어 Word 변환·저장이 모두 실패)"""
    pth = Path(path)
    if not pth.exists():
        return str(pth)
    try:
        with open(pth, "ab"):
            pass
        return str(pth)
    except (PermissionError, OSError):
        alt = pth.with_name(f"{pth.stem}_{time.strftime('%m%d-%H%M')}{pth.suffix}")
        print(f"  [안내] '{pth.name}' 이(가) 다른 프로그램에 열려 있어 '{alt.name}' 로 저장합니다.", flush=True)
        return str(alt)


def page_range_of(full: str):
    nums = [m.group(1) for m in PAGE_RE.finditer(full or "")]
    nums = [int(n) for n in nums if n.isdigit()]
    if len(nums) >= 1:
        return (nums[0], max(nums))
    return None


def save_output(translated_chunks: list, out_path, fmt: str, assets: dict = None, citation: str = "") -> list:
    """out_path: 입력 파일 경로(신규 규칙으로 이름 생성) 또는 출력 경로(하위 호환).
    저장된 파일 경로 리스트를 반환."""
    full, notes = assemble_full(translated_chunks)
    if isinstance(out_path, dict):
        paths = out_path
    else:
        ext = Path(out_path).suffix.lower().lstrip(".")
        paths = {f: str(Path(out_path).with_suffix("." + f)) for f in ({"docx+pdf": ["docx", "pdf"]}.get(fmt, [fmt]))} \
            if ext in ("docx", "pdf", "md") else output_paths(out_path, fmt, page_range_of(full))
    paths = {k: _writable_path(v) for k, v in paths.items()}
    saved = []
    if "md" in paths:
        save_md(None, paths["md"], assets, full, notes, citation=citation); saved.append(paths["md"])
    docx_path = paths.get("docx")
    if "docx" in paths or "pdf" in paths:
        import tempfile
        tmp_docx = docx_path or os.path.join(tempfile.gettempdir(), f"tr_render_{os.getpid()}.docx")
        n_fn = save_docx(None, tmp_docx, assets, full, notes, citation=citation)
        if docx_path:
            saved.append(docx_path)
        if "pdf" in paths:
            from doc_render import docx_to_pdf
            engine = ""
            try:
                engine = docx_to_pdf(tmp_docx, paths["pdf"])
            except Exception as e:
                print(f"    (DOCX->PDF 변환 오류: {e})", flush=True)
            try:
                if engine:
                    print(f"  PDF 변환: {engine} 사용 (각주 {n_fn}개, 페이지 하단)", flush=True)
                else:
                    print("  PDF 변환: Word/LibreOffice 를 쓸 수 없어 내장 렌더러 사용 (각주는 문서 끝 '원주')", flush=True)
                    save_pdf_reportlab(None, paths["pdf"], assets, full, notes, citation=citation)
                saved.append(paths["pdf"])
            finally:
                if not docx_path:
                    try:
                        Path(tmp_docx).unlink()
                    except Exception:
                        pass
    try:
        (BASE_DIR / "last_output.md").write_text(full + "\n\n" + "\n".join(f"[^{k}]: {v}" for k, v in notes.items()),
                                                 encoding="utf-8")
    except Exception:
        pass
    return saved


# ──────────────────────────────────────────────
# QA (지침 6절 / 부록 A): 쪽 마커 순차, 각주 짝, 한글 글자수/원문 단어수 비율
# ──────────────────────────────────────────────
_HANGUL_RE = re.compile(r"[가-힣]")
RATIO_LOW  = 1.4      # 정상 약 1.8~2.4 (지침 부록 A). 이보다 낮으면 누락/요약 의심
RATIO_HIGH = 3.4      # 이보다 높으면 쪽 마커 밀림/중복 의심


def hangul_count(text: str) -> int:
    return len(_HANGUL_RE.findall(text or ""))


def src_word_count(text: str) -> int:
    t = PAGE_RE.sub(" ", MARKER_RE.sub(" ", text or ""))
    t = FNREF_RE.sub("", t)
    return len(re.findall(r"[A-Za-z][A-Za-z'\-]*", t))


def looks_summarized(src: str, tr: str) -> bool:
    """원문 대비 번역이 지나치게 짧으면(요약본/거부 의심) True."""
    w = src_word_count(src)
    if w < 120:
        return False
    return hangul_count(tr) / w < 1.2


def run_qa(full: str, notes: dict, stats: dict = None, vision: bool = False) -> list:
    stats = stats or {}
    lines = []
    warn = 0
    # 1) 쪽 마커
    toks = [m.group(1) for m in PAGE_RE.finditer(full)]
    nums = [int(t) for t in toks if t.isdigit()]
    if PAGE_MARKERS:
        if not nums:
            lines.append("[쪽 마커] 없음 — 원서 쪽번호를 찾지 못함 (bat: Options > 4 Page markers / GUI: [첫 쪽 번호] 에서 지정 가능)")
            warn += 1
        else:
            exp = set(range(nums[0], max(nums) + 1))
            missing = sorted(exp - set(nums))
            backs = [(a, b) for a, b in zip(nums, nums[1:]) if b <= a]
            src = stats.get("page_source") or ("vision" if vision else "")
            lines.append(f"[쪽 마커] p.{nums[0]}–{max(nums)} ({len(nums)}개, 출처: {src or '-'})")
            if missing:
                lines.append(f"  ! 빠진 쪽: {missing[:30]}{' ...' if len(missing) > 30 else ''}"); warn += 1
            if backs:
                lines.append(f"  ! 역순/중복: {backs[:10]}"); warn += 1
    # 2) 각주
    refs = FNREF_RE.findall(full)
    if notes or refs or FOOTNOTE_MODE != "strip":
        orphan_defs = [k for k in notes if k not in set(refs)]
        lines.append(f"[각주] 모드={FOOTNOTE_MODE}, 본문 참조 {len(refs)}개, 정의 {len(notes)}개"
                     + (f", 원문 각주 {len(stats.get('notes') or [])}개" if stats.get("notes") else ""))
        if orphan_defs:
            lines.append(f"  ! 참조 없는 각주 정의 {len(orphan_defs)}개 (문서에 미포함): {orphan_defs[:10]}"); warn += 1
    # 3) 쪽별 분량 비율
    segs = {}
    parts = PAGE_RE.split(full)
    # split 결과: [앞부분, n1, 텍스트1, n2, 텍스트2, ...]
    for i in range(1, len(parts) - 1, 2):
        segs[parts[i]] = segs.get(parts[i], "") + parts[i + 1]
    pw = stats.get("page_words") or {}
    flagged = []
    # 3-0) 내용 없는 쪽: 마커만 있고 번역문이 없음 (텍스트 레이어 없는 쪽을 놓친 경우 등)
    en_pages = {str(x) for x in (stats.get("endnote_pages") or [])}
    keys = list(segs)
    empty = [n for n in keys
             if hangul_count(segs[n]) == 0 and not MARKER_RE.search(segs[n]) and n not in en_pages]
    if empty:
        warn += 1
        lines.append(f"  ! 번역문이 없는 쪽 {len(empty)}개: {empty[:30]}{' ...' if len(empty) > 30 else ''}"
                     " (장 끝 미주·그림 전용 쪽이면 정상)")
    if stats.get("hybrid"):
        lines.append(f"[혼합 PDF] 텍스트 레이어 없는 {stats.get('vision_page_count', 0)}쪽은 이미지(비전)로 번역")
    if stats.get("endnote_pages"):
        ep = stats["endnote_pages"]
        how = {"strip": "생략", "commentary": "논평성만 각주로", "all": "전체 각주로"}.get(FOOTNOTE_MODE, "")
        lines.append(f"[미주] 장 끝 Notes 절 p.{ep[0]}–{ep[-1]} → {how}"
                     + (f" (본문과 연결 {stats.get('endnotes_linked')}개)" if stats.get("endnotes_linked") else ""))
    if pw and segs:
        ratios = {}
        for n, txt in segs.items():
            w = pw.get(str(n))
            if w and w >= 80:
                ratios[n] = hangul_count(txt) / w
        for n, r in ratios.items():
            if r < RATIO_LOW:
                flagged.append(f"p.{n} 비율 {r:.2f} (누락/요약 의심)")
            elif r > RATIO_HIGH:
                flagged.append(f"p.{n} 비율 {r:.2f} (마커 위치 밀림 의심)")
        if ratios:
            vals = sorted(ratios.values())
            lines.append(f"[분량] 한글 글자수/원문 단어수 중앙값 {vals[len(vals)//2]:.2f} (정상 약 1.8~2.4)")
    elif segs:
        counts = {n: hangul_count(t) for n, t in segs.items()}
        vals = sorted(counts.values())
        med = vals[len(vals) // 2] if vals else 0
        keys = list(counts)
        for k, n in enumerate(keys):
            if med and counts[n] < med * 0.35 and 0 < k < len(keys) - 1:
                flagged.append(f"p.{n} 한글 {counts[n]}자 (중앙값 {med}자 대비 적음 — 그림/표 쪽이 아니면 누락 의심)")
        lines.append(f"[분량] 쪽당 한글 글자수 중앙값 {med}자")
    if flagged:
        warn += len(flagged)
        lines += ["  ! " + f for f in flagged[:40]]
    fails = re.findall(r"\[번역 실패[^\]]*\]", full)
    if fails:
        lines.append(f"  ! 번역 실패 구간 {len(fails)}개: {fails[:10]}"); warn += len(fails)
    head = f"QA 점검: {'이상 없음' if warn == 0 else f'확인 필요 {warn}건'}"
    report = [head] + lines
    try:
        (BASE_DIR / "last_qa.txt").write_text("\n".join(report), encoding="utf-8")
    except Exception:
        pass
    return report


def finalize_and_save(translated: list, input_file: str, fmt: str, assets: dict, stats: dict = None,
                      vision: bool = False, citation: str = None) -> list:
    """저장 + QA + 정리. 저장 실패 시 MD 로라도 남긴다. 저장된 경로 리스트 반환."""
    full, notes = assemble_full(translated)
    paths = output_paths(input_file, fmt, page_range_of(full))
    saved = []
    try:
        if citation is None:
            citation = get_citation()
        saved = save_output(translated, paths, fmt, assets, citation=citation)
        for pth in saved:
            print(f"\n저장 완료: {pth}", flush=True)
    except Exception as e:
        try:
            print(f"\n[저장 오류] {e}", flush=True)
        except Exception:
            print("\n[저장 오류]", flush=True)
        try:
            md_path = output_paths(input_file, "md", page_range_of(full))["md"]
            save_md(None, md_path, assets, full, notes)
            print(f"  대신 Markdown 으로 저장: {md_path}", flush=True)
        except Exception:
            pass
        saved = []
    try:
        print(flush=True)
        for ln in run_qa(full, notes, stats, vision):
            print("  " + ln, flush=True)
    except Exception as e:
        print(f"  (QA 점검 실패: {e})", flush=True)
    if saved:
        cleanup_after_save(input_file, "md" if fmt == "md" else "docx")
    return saved


# 하위 호환 이름
save_pdf = save_pdf_reportlab


# ──────────────────────────────────────────────
# 비전(스캔 PDF) 번역 지원 — 텍스트 레이어가 없는 순수 스캔 이미지용
# ──────────────────────────────────────────────
VISION_PAGES_PER_REQ = 2

CLIP_RE = re.compile(r"\[\[\s*(?:CLIP|IMG)\s*:\s*(\d{1,3})\s*:\s*(\d{1,3})\s*[-~]\s*(\d{1,3})\s*\]\]")

VISION_PROMPT_EXTRA = """
## Scanned-page input (vision mode)
The user message contains scanned page images of the paper, each preceded by its page label ("Page N"). Read the pages and translate their content into Korean under the same contract and style rules, in reading order.

Additional rules for scanned pages:
- The topmost line of a page is usually a RUNNING HEADER: a page number plus the authors' names or an abbreviated title (e.g. "518 Myongsob Kim and Horace Jeffery Hodges", "Korea as a Clashpoint of Civilizations 523"). NEVER translate it, never output it, and NEVER merge it into a nearby sentence — a sentence interrupted by a page break must continue seamlessly with the header text simply dropped. Also skip standalone page numbers and journal/database boilerplate.
- {footnote_rule}
- {endnote_rule}
- {page_rule}
- Reconstruct paragraphs across line breaks and page breaks: if the last paragraph of one page continues on the next page (or continues from the previous request, shown in <previous_translation>), merge it into a single Korean paragraph. Never translate a half-sentence as if it were complete.
- Use markdown headings: # for the paper/chapter title (prefix the chapter number if printed, e.g. "# 3. ..."), ## for section headings, ### for sub-section headings; **bold** where the layout clearly shows it. Put the author name line right after the title as a plain line (not a heading).
- If a page contains a figure, chart, map, diagram, table, or complex display equation that cannot be faithfully rendered as text, insert a marker on its own line at that position: [[CLIP:P:T-B]] where P is the page label number and T-B are the top and bottom of that region as integer percentages of the page height (0 = top edge, 100 = bottom edge). Use EXACTLY this CLIP syntax — never [[IMG:...]], which is reserved for a different pipeline. Include the whole graphic in the range but not the body text around it. Then translate the caption as normal text right after the marker. Transcribe simple one-line equations as text instead of using a marker. Never use CLIP markers for ordinary paragraphs, block quotations, indented quotes, epigraphs, or any content that is plain text — translate those in full.
- If a page is blank or contains no translatable content, output nothing for it.
"""


VISION_FN_STRIP = ("Skip footnotes (smaller text at the bottom of a page) entirely, and remove "
                   "footnote callout numbers from the body text.")
VISION_FN_NOTES = (
    "FOOTNOTES ({which}): where the body text carries a footnote callout (superscript number N), "
    "write [^N] attached directly to the corresponding Korean word or sentence. After the body text of the "
    "LAST page in this request, output the footnote definitions, one per line: `[^N]: (원주 N) 한국어 번역`. "
    "Translate commentary fully into Korean but keep bibliographic citations (authors, titles, journals, "
    "publishers, years, pages) in their original English. {drop_rule} A footnote that continues from the "
    "previous page belongs to the note it continues; never translate footnote text into the body.")
VISION_DROP_COMMENTARY = ("Include ONLY notes that contain explanation, argument, data, or discussion; for a "
                          "purely bibliographic note (only citations/'Ibid.'/'See ...'), omit both its definition "
                          "and its [^N] callout.")
VISION_EN_STRIP = ("ENDNOTES: a 'Notes' section at the end of a chapter (a numbered list of citations/notes) "
                   "must be skipped entirely, including its heading; remove endnote callout numbers from the body.")
VISION_EN_NOTES = (
    "ENDNOTES: body callouts that refer to a 'Notes' section at the end of the chapter are written [^eN] "
    "(N = the printed number), attached directly to the Korean word or sentence. When you reach the chapter's "
    "'Notes' section, do not output its heading as body text; instead output each entry as one line "
    "`[^eN]: (원주 N) 한국어 번역` (commentary in Korean, citations in original English). {drop_rule}")
VISION_PAGE_RULE = (
    "PAGE MARKERS: at the point where each page of the original book/journal begins, write [[PAGE:N]] on its "
    "own line, where N is the page number PRINTED on that page (in the running header or footer; if it is "
    "not visible, infer it from neighbouring pages). If one image shows two book pages side by side, emit a "
    "marker for each page. If a paragraph continues across a page break, put the marker inline inside the "
    "Korean paragraph at the corresponding point instead of breaking the paragraph.")


def build_vision_system_blocks(glossary: str = "") -> list:
    if FOOTNOTE_MODE == "strip":
        fn_rule = VISION_FN_STRIP
    elif FOOTNOTE_MODE == "all":
        fn_rule = VISION_FN_NOTES.format(which="translate all notes", drop_rule="")
    else:
        fn_rule = VISION_FN_NOTES.format(which="commentary notes only", drop_rule=VISION_DROP_COMMENTARY)
    if FOOTNOTE_MODE == "strip":
        en_rule = VISION_EN_STRIP
    else:
        en_rule = VISION_EN_NOTES.format(drop_rule="" if FOOTNOTE_MODE == "all" else
                                         "Omit purely bibliographic entries and their callouts.")
    page_rule = VISION_PAGE_RULE if PAGE_MARKERS else "Do not output page numbers."
    glossary = merge_glossary(glossary)
    prompt = SYSTEM_PROMPT + VISION_PROMPT_EXTRA.format(footnote_rule=fn_rule, endnote_rule=en_rule,
                                                        page_rule=page_rule)
    blocks = [{"type": "text", "text": prompt}]
    if glossary and glossary.strip():
        blocks.append({"type": "text",
                       "text": "## Glossary (use these renderings consistently)\n" + glossary.strip()})
    blocks[-1]["cache_control"] = {"type": "ephemeral"}
    return blocks


def _image_block(path: str) -> dict:
    import base64
    data = base64.standard_b64encode(Path(path).read_bytes()).decode("ascii")
    ext = Path(path).suffix.lower()
    media = "image/png" if ext == ".png" else "image/jpeg"
    return {"type": "image", "source": {"type": "base64", "media_type": media, "data": data}}


def build_vision_user_content(group: list, total_pages: int, prev_tr: str = "") -> list:
    """group: [{'label': n, 'path': ...}, ...] -> content 블록 리스트."""
    content = []
    if prev_tr:
        content.append({"type": "text",
                        "text": "<previous_translation>\n" + prev_tr.strip() + "\n</previous_translation>"})
    for g in group:
        if g.get("printed"):
            content.append({"type": "text",
                            "text": f"Page image — original page number {g['label']} "
                                    f"(use [[PAGE:{g['label']}]] for this page):"})
        else:
            content.append({"type": "text", "text": f"Page {g['label']} (of {total_pages}):"})
        content.append(_image_block(g["path"]))
    content.append({"type": "text",
                    "text": "Translate the page(s) above into Korean now, following the system rules. "
                            "Output only the translation."})
    return content


def build_glossary_vision(client, model: str, vision_pages: list, max_pages: int = 12) -> tuple:
    """스캔 PDF용 용어집: 앞쪽 페이지 이미지들로 1회 호출. (glossary, usage) 반환."""
    content = []
    for g in vision_pages[:max_pages]:
        content.append({"type": "text", "text": f"Page {g['label']}:"})
        content.append(_image_block(g["path"]))
    content.append({"type": "text", "text": "Build the glossary from these scanned pages."})
    kwargs = get_thinking_kwargs(model, "low")
    resp = call_with_retry(lambda: client.messages.create(
        model=model,
        max_tokens=6000,
        system=GLOSSARY_SYSTEM,
        messages=[{"role": "user", "content": content}],
        **kwargs,
    ))
    return _parse_glossary(extract_response_text(resp.content)), resp.usage


def namespace_vision_notes(text: str, group_idx: int, group: list = None) -> str:
    """비전 출력 후처리.
    - 각주 [^N] / 미주 [^eN] 를 [^vN] 로 통일 → assemble_full 에서 '같은 번호의 가장 가까운 앞 참조'와
      위치 순서로 짝지음 (미주 정의가 다른 요청에서 나와도, 장마다 번호가 다시 시작해도 연결됨)
    - 인쇄 쪽번호를 알려준 요청인데 모델이 [[PAGE:N]] 을 빠뜨렸으면 비슷한 위치에 복원"""
    t = re.sub(r"\[\^e?(\d{1,3})\]", lambda m: f"[^v{m.group(1)}]", text or "")
    if group and all(g.get("printed") for g in group) and PAGE_MARKERS:
        present = set(PAGE_RE.findall(t))
        n = len(group)
        for k, g in enumerate(group):
            lab = str(g["label"])
            if lab in present:
                continue
            tok = f"[[PAGE:{lab}]]"
            body, _, defs = t.partition("\n[^v")          # 미주 정의 목록 앞쪽 본문만 대상
            paras = body.split("\n\n")
            pos = 0 if k == 0 else min(len(paras), round(len(paras) * k / n))
            paras.insert(pos, tok)
            t = "\n\n".join(paras) + (("\n[^v" + defs) if _ else "")
    return t


VISION_TOKEN_RE = re.compile(r"\[\[VISION:(\d+)\]\]")


def build_units(text: str, stats: dict) -> list:
    """혼합 PDF 대응: 본문 텍스트를 [[VISION:k]] 지점에서 나눠, 순서대로
    {'type': 'text', 'src': 청크} / {'type': 'vision', 'group': [쪽...]} 단위 목록을 만든다."""
    runs = stats.get("vision_runs") or []
    units = []
    pos = 0
    for m in VISION_TOKEN_RE.finditer(text):
        seg = text[pos:m.start()].strip()
        if seg:
            units += [{"type": "text", "src": c} for c in chunk_text(seg)]
        k = int(m.group(1))
        if k < len(runs):
            units += [{"type": "vision", "group": g} for g in group_vision_pages(runs[k])]
        pos = m.end()
    seg = text[pos:].strip()
    if seg:
        units += [{"type": "text", "src": c} for c in chunk_text(seg)]
    return units


def group_vision_pages(vision_pages: list, per_req: int = VISION_PAGES_PER_REQ) -> list:
    return [vision_pages[i:i + per_req] for i in range(0, len(vision_pages), per_req)]


def process_clip_markers(translated: str, pdf_path: str, vision_pages: list,
                         assets: dict, assets_dir) -> str:
    """[[CLIP:P:T-B]] 마커를 실제 클립 이미지 자산 + [[IMG:id]] 마커로 치환."""
    from pdf_extract import clip_page_region
    page_map = {g["label"]: g["pdf_page"] for g in vision_pages}
    assets_dir = Path(assets_dir)
    assets_dir.mkdir(parents=True, exist_ok=True)

    def _rep(m):
        label, top, bot = int(m.group(1)), int(m.group(2)), int(m.group(3))
        if label not in page_map or bot <= top:
            return ""
        n = sum(1 for k in assets if k.startswith("clip_")) + 1
        img_id = f"clip_p{label:03d}_{n:02d}"
        out_path = assets_dir / f"{img_id}.png"
        try:
            clip = clip_page_region(pdf_path, page_map[label], top, bot, out_path)
        except Exception:
            return ""
        assets[img_id] = {"path": str(out_path), "kind": "fig", "page": label,
                          "width_pt": round(clip.width, 1), "height_pt": round(clip.height, 1)}
        return f"\n\n[[IMG:{img_id}]]\n\n"

    out = CLIP_RE.sub(_rep, translated)
    try:
        (assets_dir / "manifest.json").write_text(
            json.dumps(assets, ensure_ascii=False, indent=2), encoding="utf-8")
    except Exception:
        pass
    return re.sub(r"\n{3,}", "\n\n", out)


# ──────────────────────────────────────────────
# 저장 성공 후 임시 산출물 정리
# ──────────────────────────────────────────────
def cleanup_after_save(input_file: str, fmt: str) -> None:
    """번역 결과 저장이 성공한 뒤 assets/<stem>/ 의 임시 파일을 정리한다.

    - DOCX/PDF: 이미지가 문서에 포함되므로 assets/<stem>/ 전체 삭제
    - MD: 본문이 이미지 파일 경로를 참조하므로 클립 이미지는 남기고
          비전 모드 페이지 렌더(pages/)만 삭제
    빈 assets/ 루트 폴더도 제거한다. 실패해도 조용히 넘어간다.
    """
    import shutil
    try:
        d = assets_dir_for(input_file)
        if fmt == "md":
            p = Path(d) / "pages"
            if p.exists():
                shutil.rmtree(p, ignore_errors=True)
        else:
            if Path(d).exists():
                shutil.rmtree(d, ignore_errors=True)
        try:
            if ASSETS_ROOT.exists() and not any(ASSETS_ROOT.iterdir()):
                ASSETS_ROOT.rmdir()
        except Exception:
            pass
    except Exception:
        pass
